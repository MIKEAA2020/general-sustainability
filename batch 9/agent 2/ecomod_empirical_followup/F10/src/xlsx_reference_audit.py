"""Full, source-specific bounded workbook retry; no map-accuracy calculation."""
from pathlib import Path,PurePosixPath
import json,csv,io,zipfile,hashlib,re,unicodedata,argparse,collections
import openpyxl,requests
from reference_audit import Profile,number,code,sha,require,acquire
ROOT=Path(__file__).resolve().parents[1]

def token(value):
    if value is None:return ''
    t=unicodedata.normalize('NFC',str(value).strip())
    if re.fullmatch(r'[+-]?\d+(?:[.,]\d+)?',t):
        v=number(t)
        if v is not None and v.is_integer():return str(int(v))
    return t

def rowhash(row,fields):return hashlib.sha256(json.dumps([token(row.get(k))for k in fields],ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
def safe_zip(z,cfg):
    infos=z.infolist();require(bool(infos),'Empty workbook');require(len({x.filename for x in infos})==len(infos),'Duplicate workbook member')
    require(all(not PurePosixPath(x.filename).is_absolute()and '..'not in PurePosixPath(x.filename).parts and not x.flag_bits&1 for x in infos),'Unsafe workbook archive')
    require(all(x.file_size<=cfg['worksheet_uncompressed_cap']for x in infos)and sum(x.file_size for x in infos)<=cfg['archive_uncompressed_total_cap'],'Workbook decompression cap exceeded');return infos

def get_workbook(path,meta,cfg,offline):
    if path.exists():require(path.stat().st_size==meta['bytes']and sha(path)==meta['sha256'],'Workbook cache drift');return
    require(not offline,'Missing workbook cache');require(meta['bytes']<=cfg['workbook_download_cap'],'Workbook transfer cap');r=requests.get(meta['url'],headers={'If-Match':meta['etag'],'Accept-Encoding':'identity'},stream=True,timeout=(20,90));total=0;h=hashlib.sha256();path.parent.mkdir(exist_ok=True,parents=True)
    try:
        require(r.status_code==200 and r.headers.get('ETag')==meta['etag'],'Workbook request/ETag drift')
        with path.open('wb')as f:
            for b in r.iter_content(262144):
                total+=len(b);require(total<=min(cfg['workbook_download_cap'],meta['bytes']),'Workbook response limit');h.update(b);f.write(b)
        require(total==meta['bytes']and h.hexdigest()==meta['sha256'],'Workbook response identity drift')
    except Exception:
        if path.exists():path.unlink()
        raise
    finally:r.close()

def run(cache,offline=False):
    frozen=json.loads((ROOT/'metadata/reference/xlsx_retry_freeze_F08.json').read_text());cfg=json.loads((ROOT/'configs/xlsx_retry_F08.json').read_text())
    for name,key in [('protocol/XLSX_RETRY_F08.md','protocol_sha256'),('configs/xlsx_retry_F08.json','configuration_sha256')]:require(sha(ROOT/name)==frozen[key],'Frozen retry rules changed')
    meta=json.loads((ROOT/'metadata/reference/xlsx_full_source_F08.json').read_text());head=frozen['xlsx_HEAD'];hh=requests.structures.CaseInsensitiveDict(head['headers']);require(meta['url']==head['url']and meta['etag']==hh['ETag']and meta['bytes']==int(hh['Content-Length']),'Frozen workbook identity drift');cache.mkdir(exist_ok=True,parents=True);path=cache/'reference.xlsx';get_workbook(path,meta,cfg,offline)
    members=[]
    with zipfile.ZipFile(path)as z:
        for x in safe_zip(z,cfg):
            h=hashlib.sha256()
            with z.open(x)as f:
                for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
            members.append({'name':x.filename,'bytes':x.file_size,'sha256':h.hexdigest()})
    workbook=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        require(len(workbook.sheetnames)==1,'Unexpected multi-sheet reference source');ws=workbook.worksheets[0];it=ws.iter_rows(values_only=True);headers=[str(x).strip().upper()for x in next(it)];require(0<len(headers)<=cfg['maximum_columns']and len(headers)==len(set(headers)),'Invalid workbook header');require(all(k in headers for k in ['TARGETID','CLASS_2015','CLASS_2023','LON','LAT']),'Required workbook fields absent')
        oldmeta=json.loads((ROOT/'metadata/reference/archive_registry_F06.json').read_text())['csv_zip'];csvpath=cache/'csv_zip.zip';acquire(oldmeta,csvpath,offline);annual=sorted(k for k in headers if re.fullmatch(r'(?:CLASS|COUNT|BORDA)_\d{4}',k));saved={};cp=None
        with zipfile.ZipFile(csvpath)as z:
            names=[n for n in z.namelist()if n.endswith('.csv')];require(len(names)==1,'CSV representation drift')
            with z.open(names[0])as raw:
                reader=csv.DictReader(io.TextIOWrapper(raw,encoding='utf-8-sig'));require(all(k in reader.fieldnames for k in annual),'Annual schema mismatch requires separate review');cp=Profile(reader.fieldnames)
                for row in reader:
                    x=number(row['LON']);y=number(row['LAT']);ident=cp.add(row,(x,y));saved[ident]=(x,y,rowhash(row,annual),token(row['CLASS_2015']),token(row['CLASS_2023']),token(row.get('AMOSTRAS')),token(row.get('AMOSTRA_AM')))
        xp=Profile(headers);comparison=collections.Counter();missingids=0;raw_coords=0;formula_count=0;blankrows=0
        for values in it:
            require(len(values)<=cfg['maximum_columns'],'Workbook width cap exceeded');require(len(values)<=len(headers),'Nonempty columns outside header');values=tuple(values)+(None,)*(len(headers)-len(values))
            if all(x is None or x==''for x in values):blankrows+=1;continue
            row=dict(zip(headers,values));ident=code(row['TARGETID']);old=saved.get(ident);xy=old[:2]if old else None;xp.add({k:(''if v is None else v)for k,v in row.items()},xy);require(xp.n<=cfg['maximum_data_rows'],'Workbook data-row cap exceeded')
            formulas=[k for k,v in row.items()if isinstance(v,str)and v.startswith('=')];formula_count+=len(formulas);require(not any(k in annual for k in formulas),'Formula-based annual reference is not interpreted');lon=number(row['LON']);lat=number(row['LAT']);raw_coords+=int(lon is not None and lat is not None)
            if old:
                comparison['shared_ids']+=1;comparison['same_all_common_annual_payloads']+=int(rowhash(row,annual)==old[2]);comparison['same_2015_label_tokens']+=int(token(row['CLASS_2015'])==old[3]);comparison['same_2023_label_tokens']+=int(token(row['CLASS_2023'])==old[4]);comparison['same_training_role_tokens']+=int((token(row.get('AMOSTRAS')),token(row.get('AMOSTRA_AM')))==old[5:7])
                if lon is not None and lat is not None:
                    comparison['direct_coordinates_match_within_1e-8_deg']+=int(max(abs(lon-old[0]),abs(lat-old[1]))<=1e-8);comparison['transposed_coordinates_match_within_1e-8_deg']+=int(max(abs(lat-old[0]),abs(lon-old[1]))<=1e-8)
            else:missingids+=1
            if xp.n%20000==0:print('Audited workbook rows',xp.n,flush=True)
        summary=xp.summary();summary['coordinate_rule']='Local coverage uses TARGETID join to F06 CSV positions, previously verified against WGS84 vector geometry; workbook attributes compared separately, not silently swapped.'
        result={'revision':'F08','scope':'Full workbook reference-artifact audit; not map accuracy or independent validation','full_workbook_download_success':True,'all_data_rows_streamed':True,'workbook_sha256':meta['sha256'],'sheet_names':workbook.sheetnames,'xlsx_profile':summary,'exact_CSV_comparison':dict(comparison),'common_annual_fields_compared':annual,'IDs_not_in_CSV':missingids,'CSV_ids_not_in_workbook':len(cp.ids-xp.ids),'raw_workbook_coordinate_complete_records':raw_coords,'formula_cells_not_executed':formula_count,'fully_empty_rows_skipped':blankrows,'local_paired_reference_validation_established':False,'map_reference_class_confusion_computed':False,'national_admission_changed':False,'scope_specific_download_exception_bytes':cfg['workbook_download_cap'],'raster_limits_unchanged':True}
    finally:workbook.close()
    (ROOT/'data/reference/xlsx_audit_F08.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');(ROOT/'metadata/reference/xlsx_member_hashes_F08.json').write_text(json.dumps(members,indent=2)+'\n');print(json.dumps({k:result[k]for k in ['full_workbook_download_success','all_data_rows_streamed','exact_CSV_comparison','IDs_not_in_CSV','formula_cells_not_executed']},indent=2));print('Endpoint completeness:',summary['endpoint_nonblank_records'],'local:',summary['local_endpoint_nonblank_records'],flush=True);return result

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--cache',type=Path,required=True);p.add_argument('--offline',action='store_true');a=p.parse_args();run(a.cache,a.offline)
