"""Stage 1: reproducible access/schema and account-eligibility audit. No trends.
python src/preflight.py --cache /path/to/raw/cache [--download]
Raw inputs are external cache files, not redistributed with the package.
"""
from pathlib import Path
import argparse,csv,json,hashlib,math,collections
import requests
import numpy as np
from openpyxl import load_workbook
from netCDF4 import Dataset,num2date
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
ROOT=Path(__file__).resolve().parents[1]
INPUTS={
 'NEFBA_2026.xlsx':'https://footprint.info.yorku.ca/files/2026/04/NEFBA_Data_global_national_ecological_footprint_biocapacity.xlsx?x92789',
 'cci_50km_agb.nc':'https://data.cci.ceda.ac.uk/thredds/fileServer/esacci/biomass/data/agb/maps/v7.0/netcdf/ESACCI-BIOMASS-L4-AGB-MERGED-50000m-fv7.0.nc',
 'cci_50km_change_2020_2010.nc':'https://data.cci.ceda.ac.uk/thredds/fileServer/esacci/biomass/data/agb/maps/v7.0/netcdf/ESACCI-BIOMASS-L4-DIFF-50000m-2020-2010-fv7.0.nc'}

def writecsv(name,records):
    with (ROOT/name).open('w',newline='')as f:
        w=csv.DictWriter(f,fieldnames=list(records[0]));w.writeheader();w.writerows(records)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--cache',type=Path,default=ROOT/'.cache/raw');ap.add_argument('--download',action='store_true');args=ap.parse_args();args.cache.mkdir(parents=True,exist_ok=True)
    cfg=json.loads((ROOT/'configs/stage1_design.json').read_text());receipts=[]
    for name,url in INPUTS.items():
        p=args.cache/name
        if args.download:
            with requests.get(url,stream=True,timeout=60)as r:
                r.raise_for_status();data=bytearray()
                for chunk in r.iter_content(65536):
                    data.extend(chunk)
                    if len(data)>cfg['collection_limit_bytes_per_file']:raise ValueError('Declared size cap exceeded')
                p.write_bytes(data)
        data=p.read_bytes()
        receipts.append(dict(id=name,url=url,bytes=len(data),sha256=hashlib.sha256(data).hexdigest(),audit_date='2026-10-29',raw_redistributed=False))
    # Fail closed if raw source bytes change after the first frozen run.
    registry=ROOT/'metadata/source_registry.json'
    if registry.exists():
        old=json.loads(registry.read_text())
        assert {x['id']:x['sha256']for x in old}=={x['id']:x['sha256']for x in receipts},'Source changed: amendment/new registry required'
    registry.write_text(json.dumps(receipts,indent=2)+'\n')
    book=load_workbook(args.cache/'NEFBA_2026.xlsx',read_only=True,data_only=True)
    s=book['national_data'];it=s.iter_rows(values_only=True);header=next(it);allrows=[dict(zip(header,r))for r in it if r[0]is not None]
    keys=[(r['country_ISO3'],r['year'])for r in allrows];assert len(keys)==len(set(keys)),'Duplicate country-years'
    groups=collections.defaultdict(dict)
    for r in allrows:groups[r['country_ISO3']][r['year']]=r
    required=['biocap_'+c+'_gha'for c in cfg['account_components']]+['biocap_total_gha','area_forest_ha','area_crop_ha','pop_1000_persons']
    isnum=lambda x:isinstance(x,(int,float))and math.isfinite(x)
    audit=[];max_error=0.;mismatch_count=0;flagfreq=collections.Counter();quality_rows=[]
    for iso,years in sorted(groups.items()):
        ref=next(iter(years.values()));a,b=cfg['primary_window'];sel=[years[y]for y in range(a,b+1)if y in years]
        reasons=[]
        if len(sel)!=b-a+1:reasons.append('missing_window_year')
        flags=sorted({str(r['flag'])for r in sel});badflags=[r for r in sel if r['flag']not in cfg['account_allowed_flags']]
        if badflags:reasons.append('unpublished_or_estimated_flag')
        missing=sum(not isnum(r.get(k))for r in sel for k in required)
        if missing:reasons.append('missing_required_account_values')
        if any(isnum(r.get(k))and r[k]<0 for r in sel for k in required):reasons.append('negative_required_value')
        if any(isnum(r.get('pop_1000_persons'))and r['pop_1000_persons']<=0 for r in sel):reasons.append('nonpositive_population')
        errors=[]
        for r in sel:
            flagfreq[str(r['flag'])]+=1
            cc=['biocap_'+c+'_gha'for c in cfg['account_components']]
            if all(isnum(r[k])for k in cc+['biocap_total_gha']):
                err=abs(sum(r[k]for k in cc)-r['biocap_total_gha']);errors.append(err);max_error=max(max_error,err)
                if err>cfg['nfa_component_rounding_tolerance_gha']:mismatch_count+=1
        if errors and max(errors)>cfg['nfa_component_rounding_tolerance_gha']:reasons.append('component_total_inconsistency')
        audit.append(dict(ISO3=iso,country=ref['country_name'],UN_region=ref['UN_region'],UN_subregion=ref['UN_subregion'],first_available_year=min(years),last_available_year=max(years),window_years_present=len(sel),window_flags=';'.join(flags),missing_required_cells=missing,max_component_rounding_error_gha=max(errors)if errors else '',account_preflight_pass=not reasons,exclusion_reasons=';'.join(reasons),pilot_candidate=iso in cfg['pilot_candidates_iso3'],physical_alignment_gate='PENDING',full_study_eligible=False))
    writecsv('data/derived/account_eligibility.csv',audit)
    pilots=[next(r for r in audit if r['ISO3']==iso)for iso in cfg['pilot_candidates_iso3']];writecsv('data/derived/pilot_candidate_audit.csv',pilots)
    # Store source column definitions/flags: publication scope is not an uncertainty model.
    definitions=[list(r)for r in book['columns'].iter_rows(values_only=True)]
    (ROOT/'metadata/account_definitions.json').write_text(json.dumps(definitions,indent=2)+'\n')
    nc=[]
    for name in ['cci_50km_agb.nc','cci_50km_change_2020_2010.nc']:
        with Dataset(args.cache/name)as d:
            years=sorted({x.year for x in num2date(d['time'][:],d['time'].units)})
            missing_bounds={k:v.bounds for k,v in d.variables.items()if 'bounds'in v.ncattrs()and v.bounds not in d.variables}
            # Metadata and coordinates only. No AGB/change pixel values or signs inspected.
            rec=dict(source_id=name,dimensions={k:len(v)for k,v in d.dimensions.items()},years=years,
                globals={k:str(d.getncattr(k))for k in d.ncattrs()},
                variables={k:dict(shape=list(v.shape),dtype=str(v.dtype),attributes={x:str(v.getncattr(x))for x in v.ncattrs()})for k,v in d.variables.items()},
                first_lon_bounds=d['lon_bnds'][0].tolist(),first_lat_bounds=d['lat_bnds'][0].tolist(),
                missing_bounds_references=missing_bounds,missing_flag_semantics=('diff_qf'in d.variables and 'flag_meanings'not in d['diff_qf'].ncattrs()),
                physical_values_read=False,role='Schema/access smoke test only, not empirical national biomass results')
            nc.append(rec)
    (ROOT/'metadata/cci_schema_audit.json').write_text(json.dumps(nc,indent=2)+'\n')
    summary=dict(status='STAGE1_ACCOUNT_AND_ACCESS_PREFLIGHT_ONLY',national_rows=len(allrows),national_columns=len(header),country_identifiers=len(groups),minimum_year=min(r['year']for r in allrows),maximum_year=max(r['year']for r in allrows),primary_window=cfg['primary_window'],window_flag_counts=dict(flagfreq),account_preflight_pass=sum(r['account_preflight_pass']for r in audit),account_preflight_fail=sum(not r['account_preflight_pass']for r in audit),max_component_rounding_error_gha=max_error,rows_above_rounding_tolerance=mismatch_count,pilot_candidates=len(pilots),pilot_account_preflight_pass=sum(r['account_preflight_pass']for r in pilots),full_study_eligible_count=0,full_study_eligibility_evaluated=False,account_or_biomass_change_signs_computed=False,cci_coordinate_note='50 km label is nominal; coordinate spacing is 0.5 degrees, not equal-area',cci_change_metadata_issues={'missing_time_bounds_target':nc[1]['missing_bounds_references'],'quality_flag_meanings_absent_in_file':nc[1]['missing_flag_semantics']},no_trend_inference=True)
    (ROOT/'results/preflight_summary.json').write_text(json.dumps(summary,indent=2)+'\n')
    counts=collections.Counter()
    for r in audit:counts[(r['UN_region'],'pass'if r['account_preflight_pass']else 'fail')]+=1
    regions=sorted({r['UN_region']for r in audit});fig,ax=plt.subplots(figsize=(8,3.7),constrained_layout=True)
    x=np.arange(len(regions));passed=[counts[(r,'pass')]for r in regions];failed=[counts[(r,'fail')]for r in regions]
    ax.bar(x,passed,label='Account preflight pass',color='#377d87');ax.bar(x,failed,bottom=passed,label='Account preflight fail',color='#d6a64c');ax.set_xticks(x,regions);ax.set_ylabel('Country/territory identifiers');ax.set_title('2015–2023 account availability only — not physical-data eligibility');ax.legend(fontsize=8);fig.supxlabel('Source: Lo et al., NEFBA 2026 v1.0; derived availability counts — CC BY-SA 4.0',fontsize=7);fig.savefig(ROOT/'results/account_coverage.pdf');plt.close(fig)
    # Generated text fragments used by the implementation manuscript.
    (ROOT/'results/preflight_numbers.tex').write_text(''.join('\\newcommand{\\'+k+'}{'+str(v)+'}\n'for k,v in [('NationalRows',len(allrows)),('CountryIDs',len(groups)),('AccountPass',summary['account_preflight_pass']),('AccountFail',summary['account_preflight_fail']),('PilotPass',summary['pilot_account_preflight_pass'])]))
    print(json.dumps(summary,indent=2))

if __name__=='__main__':main()
