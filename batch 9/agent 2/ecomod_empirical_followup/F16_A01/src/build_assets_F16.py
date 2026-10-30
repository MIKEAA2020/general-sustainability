#!/usr/bin/env python3
"""Build source-pinned F16 graphics and additively extend the F15 workbook."""
from __future__ import annotations
import csv
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"


def require(ok: bool, message: str) -> None:
    if not ok:
        raise ValueError(message)


def rows(name: str) -> list[dict]:
    with (DATA / name).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize_zip(path: Path) -> None:
    """Normalize container metadata only; workbook XML and cell values are untouched."""
    with zipfile.ZipFile(path, "r") as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    buffer = BytesIO()
    fixed_modified = b"2026-10-30T00:00:00Z"
    with zipfile.ZipFile(buffer, "w") as target:
        for info, payload in entries:
            if info.filename == "docProps/core.xml":
                payload = re.sub(rb"(<dcterms:modified[^>]*>)[^<]+(</dcterms:modified>)", rb"\g<1>" + fixed_modified + rb"\g<2>", payload)
            normalized = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            normalized.compress_type = info.compress_type
            normalized.comment = info.comment
            normalized.extra = info.extra
            normalized.internal_attr = info.internal_attr
            normalized.external_attr = info.external_attr
            normalized.create_system = info.create_system
            target.writestr(normalized, payload)
    path.write_bytes(buffer.getvalue())


def save_figure(fig: plt.Figure) -> None:
    metadata = {"CreationDate": None, "ModDate": None}
    fig.savefig(ROOT / "figures/national_physical_modules_F16.pdf", bbox_inches="tight", metadata=metadata)
    fig.savefig(ROOT / "evidence/national_physical_modules_F16.png", bbox_inches="tight", dpi=180, metadata={"Software": "Matplotlib"})
    plt.close(fig)


def build_figure() -> None:
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 10.5, "pdf.fonttype": 42,
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.titleweight": "bold", "axes.titlesize": 11.5,
    })
    cmr_a = rows("CMR_reported_forest_area_2000_2020_F16.csv")
    cmr_f = rows("CMR_reported_net_forest_flux_2000_2020_F16.csv")
    ab = rows("AUS_ABARES_adjusted_forest_area_1989_2021_F16.csv")
    n21 = rows("AUS_NIR2021_stock_gains_losses_1989_2021_F16.csv")
    n23 = rows("AUS_NIR2023_forest_land_1989_2023_F16.csv")

    fig = plt.figure(figsize=(10.2, 7.0))
    grid = fig.add_gridspec(2, 2, height_ratios=[1, 1.18], hspace=0.48, wspace=0.28)
    ax_a = fig.add_subplot(grid[0, 0])
    ax_f = fig.add_subplot(grid[0, 1])
    ax_u = fig.add_subplot(grid[1, :])

    years = np.array([int(r["year"]) for r in cmr_a])
    area = np.array([int(r["reported_total_forest_ha"]) for r in cmr_a]) / 1e6
    ax_a.plot(years, area, color="#167f7b", lw=2.2)
    ax_a.scatter([2015, 2020], [area[15], area[20]], color=["#203e53", "#b44f3b"], s=35, zorder=3)
    ax_a.annotate("2015: 31.332", (2015, area[15]), xytext=(2007.7, 31.375), arrowprops={"arrowstyle": "-", "color": "#687b83"}, fontsize=9.5)
    ax_a.annotate("2020: 31.184\n−0.149 Mha", (2020, area[20]), xytext=(2011.1, 31.19), arrowprops={"arrowstyle": "-", "color": "#687b83"}, fontsize=9.5)
    ax_a.set(title="(a) Cameroon: reported forest area", ylabel="Million ha", xlim=(1999.4, 2020.6), ylim=(31.14, 31.48))
    ax_a.grid(axis="y", alpha=0.16)
    ax_a.set_axisbelow(True)

    flux = np.array([float(r["reported_net_forest_flux_tCO2e_per_year"]) for r in cmr_f]) / 1e6
    ax_f.plot(years, flux, color="#203e53", lw=2.2)
    ax_f.scatter([2015, 2020], [flux[15], flux[20]], color=["#167f7b", "#b44f3b"], s=35, zorder=3)
    ax_f.annotate("2015: −113.383", (2015, flux[15]), xytext=(2005.7, -110.8), arrowprops={"arrowstyle": "-", "color": "#687b83"}, fontsize=9.5)
    ax_f.annotate("2020: −109.600\n+3.783 MtCO₂e/yr", (2020, flux[20]), xytext=(2008.7, -126.7), arrowprops={"arrowstyle": "-", "color": "#687b83"}, fontsize=9.5)
    ax_f.set(title="(b) Cameroon: reported net forest CO₂e flux", ylabel="MtCO₂e per year", xlim=(1999.4, 2020.6), ylim=(-136, -103))
    ax_f.grid(axis="y", alpha=0.16)
    ax_f.set_axisbelow(True)

    x_ab = np.array([int(r["year_ending_june"]) for r in ab])
    y_ab = np.array([float(r["abares_adjusted_forest_area_Mha"]) for r in ab])
    x_21 = np.array([int(r["year_ending_june"]) for r in n21])
    y_21 = np.array([float(r["nir2021_forest_land_stock_Mha"]) for r in n21])
    x_23 = np.array([int(r["year_ending_june"]) for r in n23])
    y_23 = np.array([int(r["forest_land_kha"]) for r in n23]) / 1000
    ax_u.plot(x_ab, y_ab, color="#167f7b", lw=2.1, label="ABARES-adjusted area (to 2021)")
    ax_u.plot(x_21, y_21, color="#7b8890", lw=1.8, ls="--", label="NIR2021 transition stock (to 2021)")
    observed = x_23 <= 2022
    ax_u.plot(x_23[observed], y_23[observed], color="#203e53", lw=2.1, label="NIR2023 forest land (annual total)")
    ax_u.plot(x_23[-2:], y_23[-2:], color="#203e53", lw=2.1, ls=":", label="NIR2023 final year held constant")
    ax_u.scatter([2023], [y_23[-1]], facecolors="white", edgecolors="#b44f3b", s=55, linewidths=1.8, zorder=4)
    ax_u.annotate("2022–23 = 2021–22\n(carry-forward, not a new observation)", (2023, y_23[-1]), xytext=(2007.2, 138.05), arrowprops={"arrowstyle": "-", "color": "#687b83"}, fontsize=9.5)
    ax_u.set(title="(c) Australia: source- and vintage-specific forest-area series", ylabel="Million ha; years ending June", xlabel="Year ending June", xlim=(1988.5, 2024.0), ylim=(129.5, 139.0))
    ax_u.grid(axis="y", alpha=0.16)
    ax_u.set_axisbelow(True)
    ax_u.legend(loc="lower right", ncol=2, frameon=False, fontsize=9.4)
    fig.align_ylabels([ax_a, ax_f])
    save_figure(fig)

    values = {
        "Cameroon_area_Mha": {str(y): float(area[y - 2000]) for y in (2015, 2020)},
        "Cameroon_net_flux_MtCO2e_per_year": {str(y): float(flux[y - 2000]) for y in (2015, 2020)},
        "Australia_2021_endpoints_Mha": {
            "ABARES_adjusted": y_ab[-1], "NIR2021_transition_stock": y_21[-1],
            "NIR2023_forest_land": float(y_23[list(x_23).index(2021)]),
        },
        "Australia_2022_23_NIR2023_Mha": float(y_23[-1]),
        "Australia_2022_23_is_carried_forward": True,
        "series_are_not_calibration_intervals_or_AGB": True,
    }
    (ROOT / "metadata/plot_values_F16.json").write_text(json.dumps(values, indent=2) + "\n")


def style_sheet(sheet, widths: dict[int, float] | None = None) -> None:
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173E52")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(index, 24)


def build_workbook() -> None:
    source = ROOT / "inputs/national_results_F15_A01.xlsx"
    workbook = load_workbook(source)
    require(len(workbook.sheetnames) == 31, "F15 parent workbook sheet count")
    guide = workbook.create_sheet("F16 reader guide", 0)
    guide_rows = [
        ["F16 additive empirical elevation", "F15's 31 sheets are retained in value; ten F16 sheets are appended/inserted without overwriting prior results."],
        ["Cameroon area", "Reported total forest area: 31,332,386 ha (2015) to 31,183,654 ha (2020): −148,732 ha (−0.4747%). Table 9."],
        ["Cameroon flux", "Reported net forest CO2e flux: −113,383,116.17 to −109,599,895.35 tCO2e/yr (2015–2020), difference +3,783,220.82. Tables 31/37; not AGB-stock change."],
        ["Cameroon uncertainty", "Provider MC relative uncertainty: 31.10% (2015) and 32.14% (2020). No raw draws/covariances or new confidence interval."],
        ["Australia ABARES", "Adjusted forest area: 132.51 Mha (2015) to 133.57 Mha (2021): +1.06 Mha (+0.7999%). NIR2021 changes applied to ABARES total."],
        ["Australia NIR2021", "Separate transition stock: 132.98 to 134.03 Mha (2015–2021); 2021 gross loss −216 kha, gain +487 kha, reported net +270 kha."],
        ["Australia NIR2023", "Forest land: 135,318 kha (2014–15) to 137,372 kha (2021–22): +2,054 kha (+1.5179%). 2022–23 repeats 2021–22 by documented carry-forward."],
        ["Australia snapshot", "Forests of Australia (2023) attributes: 133,564,800 forest ha; latest information currency 30 June 2021; 100-m cells, COUNT=ha; not a 2023 observation."],
        ["Rounding", "Cameroon class sums exceed reported totals by 1 ha in 2011, 2012, 2015, 2016 and 2019. Australian source residuals remain within displayed rounding units; reported fields are not altered."],
        ["Inference", "All eight calibrated national 2015–2023 forest-AGB analyses remain incomplete. These modules add reported area/flux evidence, not AGB calibration or reporting-accuracy estimates."],
        ["Main figure", "Seven main figures total; the new three-panel figure shows Cameroon area/flux and Australia's separate area lineages."],
        ["AI assistance", "Arena.ai Agent Mode was used to assist with computational methods and literature review."],
    ]
    for row in guide_rows:
        guide.append(row)
    style_sheet(guide, {1: 28, 2: 120})
    guide.auto_filter.ref = "A1:B1"
    for row in range(2, guide.max_row + 1):
        guide.row_dimensions[row].height = 43

    sheet_files = [
        ("CMR area 2000-2020", "CMR_reported_forest_area_2000_2020_F16.csv"),
        ("CMR flux 2000-2020", "CMR_reported_net_forest_flux_2000_2020_F16.csv"),
        ("AUS ABARES area 1989-2021", "AUS_ABARES_adjusted_forest_area_1989_2021_F16.csv"),
        ("AUS NIR21 stock flows", "AUS_NIR2021_stock_gains_losses_1989_2021_F16.csv"),
        ("AUS NIR23 forest land", "AUS_NIR2023_forest_land_1989_2023_F16.csv"),
        ("AUS NIR23 clearing", "AUS_NIR2023_clearing_regrowth_1990_2023_F16.csv"),
        ("AUS 2023 categories", "AUS_Forests2023_category_areas_F16.csv"),
        ("AUS 2023 states", "AUS_Forests2023_state_areas_F16.csv"),
    ]
    for sheet_name, file_name in sheet_files:
        sheet = workbook.create_sheet(sheet_name)
        data = rows(file_name)
        sheet.append(list(data[0]))
        for record in data:
            sheet.append([record[key] for key in data[0]])
        style_sheet(sheet)

    provenance = json.loads((ROOT / "metadata/source_provenance_F16.json").read_text())
    sheet = workbook.create_sheet("F16 source provenance")
    sheet.append(["Title", "Used for", "URL", "SHA-256", "Bytes"])
    for record in provenance["sources"]:
        sheet.append([record["title"], record["used_for"], record["url"], record["sha256"], record["bytes"]])
    style_sheet(sheet, {1: 40, 2: 55, 3: 95, 4: 68, 5: 16})
    require(len(workbook.sheetnames) == 41, "F16 workbook sheet count")
    output = ROOT / "national_results_F16_A01.xlsx"
    fixed_time = datetime(2026, 10, 30, 0, 0, 0, tzinfo=timezone.utc)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    workbook.save(output)
    workbook.close()
    normalize_zip(output)

    old = load_workbook(source, read_only=True, data_only=False)
    new = load_workbook(output, read_only=True, data_only=False)
    checked = 0
    try:
        for name in old.sheetnames:
            require(name in new.sheetnames, "Inherited sheet missing " + name)
            left, right = old[name], new[name]
            require(left.max_row == right.max_row and left.max_column == right.max_column, "Inherited dimensions changed " + name)
            for row_left, row_right in zip(left.values, right.values):
                for value_left, value_right in zip(row_left, row_right):
                    require(value_left == value_right, f"Inherited workbook cell changed: {name}")
                    checked += 1
        require(checked == 60336, f"Unexpected inherited cell count {checked}")
    finally:
        old.close()
        new.close()
    report = {
        "file": output.name, "sha256": sha(output), "sheets": 41,
        "inherited_sheets_unchanged": 31, "new_sheets": 10,
        "inherited_cells_checked": checked, "all_F15_values_preserved": True,
    }
    (ROOT / "metadata/workbook_QA_F16.json").write_text(json.dumps(report, indent=2) + "\n")


def run() -> None:
    build_figure()
    build_workbook()
    print("Built one three-panel source-pinned figure and a 41-sheet workbook; all 60,336 inherited cells verified unchanged.")


if __name__ == "__main__":
    run()
