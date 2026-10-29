"""Generate territorial identifier and admission records without account trends.
Curated interpretation notes are separately frozen in configs/. No positive
physical admission is inferred from account availability or successful access.
"""
from pathlib import Path
import argparse,csv,json,hashlib
from openpyxl import load_workbook
from physical_support import ROOT,write_csv

def run(workbook):
    registry=json.loads((ROOT/'metadata/source_registry.json').read_text());expected=next(x for x in registry if x['id']=='NEFBA_2026.xlsx')
    assert hashlib.sha256(workbook.read_bytes()).hexdigest()==expected['sha256'],'Account source changed'
    cfg=json.loads((ROOT/'configs/territory_crosswalk_F04.json').read_text());w=load_workbook(workbook,read_only=True,data_only=True);it=w['national_data'].iter_rows(values_only=True);h=next(it);ix={v:i for i,v in enumerate(h)};inventory=[]
    for row in it:
        if row[ix['year']]==cfg['reference_year']:inventory.append({k:row[ix[k]]for k in ['country_ISO3','country_name','flag']})
    write_csv(ROOT/'data/derived/account_identifier_inventory_2015.csv',inventory)
    accounts={x['ISO3']:x for x in csv.DictReader((ROOT/'data/derived/country_availability.csv').open())};admit=[]
    for r in csv.DictReader((ROOT/'data/support/country_domain_support.csv').open()):
        admit.append(dict(ISO3=r['ISO3'],account_availability='pass'if accounts[r['ISO3']]['account_available']=='True'else'fail',boundary_id=r['boundary_id'],territorial_interpretation=cfg['territory_notes'][r['ISO3']],outside_native_domain_pct=r['outside_native_domain_pct'],endpoint_forest_mask='not_frozen',national_missingness='not_estimated',change_validation='not_established',admission='hold_not_admitted',permanent_exclusion=False))
    write_csv(ROOT/'data/support/physical_admission_F04.csv',admit)
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--workbook',type=Path,required=True);run(p.parse_args().workbook)
