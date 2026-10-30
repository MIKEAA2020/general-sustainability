#!/usr/bin/env python3
"""Create a scientific extension, original figures and a typed Excel results workbook."""
import csv,hashlib,json,re
from pathlib import Path
from html import escape
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
ROOT=Path(__file__).resolve().parents[1];FIG=ROOT/'figures';FIG.mkdir(exist_ok=True)
NAVY='#17344a';TEAL='#16847b';RED='#b85243';GREY='#687784'


def rows(name):
 with (ROOT/'data'/name).open()as f:return list(csv.DictReader(f))
def load(name):return json.loads((ROOT/'metadata'/name).read_text())
def table(headers,data):return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str,r))+' |'for r in data])

def figures():
 a=rows('accounts_candidate_changes_F12.csv');isos=['BRA','IDN','COD','CMR','CAN','FIN','FRA','AUS'];lookup={(r['ISO3'],r['domain']):float(r['percent_change'])for r in a if r['start']=='2015'and r['end']=='2023'}
 plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','savefig.facecolor':'white'})
 fig,ax=plt.subplots(figsize=(9.5,4.4));y=np.arange(8);ax.barh(y-.17,[lookup[i,'total']for i in isos],.32,color=TEAL,label='Total biocapacity account');ax.barh(y+.17,[lookup[i,'forest_component']for i in isos],.32,color=NAVY,label='Forest account component');ax.axvline(0,color=GREY,lw=.8);ax.set(yticks=y,yticklabels=isos,xlabel='Published account change, 2015–2023 (%)',title='Account composition: different signs within the same ledger');ax.invert_yaxis();ax.legend(frameon=False,loc='lower right',fontsize=9);fig.tight_layout();fig.savefig(FIG/'Figure1_account_composition_F12.png',dpi=180);plt.close(fig)
 r=rows('IDN_C41_extent_changes_F12.csv');d={x['group']:float(x['percent_change'])for x in r if x['level']=='national'and x['percent_change']!=''};s=load('IDN_collection_sensitivity_F12.json')
 fig,axs=plt.subplots(1,2,figsize=(10,4.2),gridspec_kw={'width_ratios':[1.15,1]});names=['Forest formation','Mangrove','Peat swamp forest','Combined mapped forest'];vals=[d[x]for x in ['forest_formation','mangrove','peat_swamp_forest','mapped_forest']];axs[0].barh(names,vals,color=[NAVY,TEAL,RED,NAVY]);axs[0].invert_yaxis();axs[0].axvline(0,color=GREY,lw=.8);axs[0].set(title='Collection 4.1: legend matters',xlabel='2015–2023 mapped-area change (%)');axs[1].bar(['Collection 3','Collection 4.1'],[s['C3_forest_percent_change'],s['C41_forest_percent_change']],color=[GREY,TEAL]);axs[1].axhline(0,color=GREY,lw=.8);axs[1].set(title='Collection AND support sensitivity',ylabel='Mapped forest change (%)');axs[1].text(.5,.05,'Not a pure version effect or confidence interval',ha='center',transform=axs[1].transAxes,fontsize=8);fig.tight_layout();fig.savefig(FIG/'Figure2_Indonesia_definition_source_F12.png',dpi=180);plt.close(fig)
 r=[x for x in rows('endpoint_conversion_path_bounds_F12.csv')if x['constraint_level']=='finest_reported_geography'];labels=[]
 for x in r:labels.append(x['country']+' '+('F+S'if x['initial_group']=='formations_plus_savanna'else'F')+' → '+x['final_target'].replace('_',' '))
 fig,ax=plt.subplots(figsize=(10,4.5))
 for i,x in enumerate(r):
  lo,hi=float(x['lower_unique_endpoint_area_ha'])/1e6,float(x['upper_unique_endpoint_area_ha'])/1e6;c=TEAL if x['country']=='IDN'else NAVY;ax.plot([lo,hi],[i,i],lw=5,color=c,alpha=.8);ax.plot([lo,hi],[i,i],'|',color=c,ms=11)
 ax.set(yticks=range(len(r)),yticklabels=labels,xlabel='Unique initial-to-final mapped-class area (million ha)',title='Endpoint-conversion feasible ranges — not confidence intervals');ax.invert_yaxis();ax.set_xlim(left=-.3);fig.tight_layout();fig.savefig(FIG/'Figure3_endpoint_conversion_bounds_F12.png',dpi=180);plt.close(fig)
 j=load('FIN_cycle_volume_inference_F12.json');fig,ax=plt.subplots(figsize=(9.5,3.0))
 for y,key,color,label in [(0,'approx_95pct_simultaneous_endpoint_difference_interval_Mm3',TEAL,'Marginal-normal Bonferroni'),(1,'moment_only_95pct_two_sided_sensitivity_interval_Mm3',GREY,'Moment-bound sensitivity')]:
  lo,hi=j[key];ax.plot([lo,hi],[y,y],lw=6,color=color);ax.plot(j['point_difference_Mm3'],y,'o',color=NAVY,ms=7)
 ax.axvline(0,color=RED,lw=1);ax.set(yticks=[0,1],yticklabels=['Approximate 95% sampling interval','95% moment-bound sensitivity'],xlabel='NFI13 minus NFI12 growing-stock volume (million m³)',title='Finland: dependence is not assumed away');ax.invert_yaxis();fig.tight_layout();fig.savefig(FIG/'Figure4_Finland_sampling_assumptions_F12.png',dpi=180);plt.close(fig)


def manuscript():
 a=rows('accounts_candidate_changes_F12.csv');isos=['BRA','IDN','COD','CMR','CAN','FIN','FRA','AUS'];d={(r['ISO3'],r['domain']):float(r['percent_change'])for r in a if r['start']=='2015'and r['end']=='2023'}
 account_table=table(['ISO3','Total','− fishing','− built','− both','Forest account','Per person'],[[i]+[f'{d[i,k]:+.2f}'for k in ['total','terrestrial_excluding_fishing','excluding_built','terrestrial_excluding_built','forest_component','per_person_total']]for i in isos])
 p=[r for r in rows('endpoint_conversion_path_bounds_F12.csv')if r['constraint_level']=='finest_reported_geography'];path_table=table(['Country / initial group','2023 target','Lower (Mha)','Upper (Mha)'],[[r['country']+' / '+('F+S'if r['initial_group']=='formations_plus_savanna'else'F'),r['final_target'].replace('_',' '),f"{float(r['lower_unique_endpoint_area_ha'])/1e6:.4f}",f"{float(r['upper_unique_endpoint_area_ha'])/1e6:.4f}"]for r in p])
 fin=load('FIN_cycle_volume_inference_F12.json');lo,hi=fin['approx_95pct_simultaneous_endpoint_difference_interval_Mm3'];mlo,mhi=fin['moment_only_95pct_two_sided_sensitivity_interval_Mm3'];refs=load('BRA_reference_compatibility_F12.json')['diagnostics'];regional=rows('national_mapped_regional_compensation_F12.csv');idr=next(x for x in regional if x['ISO3']=='IDN')
 text=f'''# National account composition, mapped forest extent and inventory-cycle stock change

## An empirical extension with conditional conversion bounds

Amin Abaee · Independent Researcher, Tehran, Iran · ORCID: https://orcid.org/0000-0002-0019-1842

30 October 2026 · F12_A01 · Additive empirical extension; the earlier manuscript remains unchanged.

## Abstract

An aggregate biocapacity account and a physical ecological asset need not move together, but an empirical comparison must identify its quantities, temporal operators and uncertainty. We compute 2015–2023 account changes for all 183 eligible identifiers in the 2026 national accounts, complete national MapBiomas extent and transition analyses for Brazil and Indonesia, and official Finnish inventory-cycle biomass and volume summaries. Total biocapacity increases in 91 identifiers and decreases in 92; 68 of the increasing totals coexist with a declining forest account component. Indonesia combines a 5.46% total-account increase with a 0.66% decline in Collection 4.1 mapped forest extent. Brazil instead has both a decreasing total account and decreasing mapped formations. A layered marked-flow model constrains unique endpoint conversions without equating repeated annual events with unique hectares. Its geographically refined Indonesia forest-to-oil-palm range is approximately 0.112–2.326 million hectares, conditional on the published classifications. Finnish growing-stock volume increases by 76.772 million cubic metres between NFI12 and NFI13; an approximate 95% Bonferroni interval remains positive without assuming cross-inventory independence. These are substantive national results, but neither mapped extent nor inventory-cycle volume establishes a calibrated 2015–2023 national above-ground biomass change. Reporting should preserve those evidential distinctions rather than erase the new evidence or promote it to a different target.

Keywords: Biocapacity; natural capital accounting; forest extent; national forest inventory; land-use change; partial identification; uncertainty.

## 1. Questions and estimands

The empirical question is whether an improving account can coexist with loss in a specified ecological asset, and what the evidence establishes about conversion and spatial compensation. Three quantities are kept distinct: published biocapacity in global hectares (gha); areas assigned to operational mapped land-cover classes in hectares; and inventory growing-stock volume in cubic metres or biomass dry mass in tonnes. None is a universal index of ecological condition. A forest plantation, natural forest, temporarily unstocked forest and an annual mapped formation may belong to different reporting categories.

The primary national above-ground biomass target retains its 2015 and 2023 endpoints and common-support requirements. National mapped-cover comparisons and Finnish inventory-cycle inference are additional, explicitly typed analyses, not replacements for that target. The eight inherited candidates remain Brazil, Indonesia, the Democratic Republic of the Congo, Cameroon, Canada, Finland, France and Australia. They are a logistical case set, not a probability sample. Source feasibility and prior country accounts or published trend headlines were known during this exploratory extension; no untouched held-out or outcome-naive selection is claimed.

## 2. Data and methods

### 2.1. Accounts, eligibility and composition

The 2026 national workbook supplies 13,652 country-year rows and 236 identifiers. The inherited availability rules admit 183 identifiers with all annual records in 2015–2023, allowed publication flags, positive population, nonnegative required quantities and numerical component consistency. We calculate actual endpoint changes, all five component contributions, total-minus-fishing, total-minus-built, total-minus-both and per-person sensitivities. The inherited 2010–2020 secondary window is also executed for 182 comparable members of the same primary cohort; South Sudan lacks an admissible 2010 endpoint. The additional 2015–2022 sensitivity does not supersede that secondary window. [1]

For component k, its contribution in percentage points of the initial total is 100(B_k,2023 − B_k,2015)/B_total,2015. A decline in the forest account component is an account result, not an independently measured forest loss. For forest gha = area × accounting factor, the symmetric decomposition is ΔB_f = ΔA(f_0+f_1)/2 + Δf(A_0+A_1)/2. The implied gha/ha factor is not interpreted as physical forest yield or condition. Component-sum tolerances account only for rounding.

### 2.2. Complete provider tables, not extrapolated local windows

Brazil uses the complete official Collection 11 biome-state workbook: 960 leaf-class geographic rows, 27 states, six biomes and an approximately 850.645-million-hectare reported footprint. Both endpoints and all eight consecutive annual transition intervals come from the same collection. The formations group F comprises Forest Formation (3), Mangrove (5), Floodable Forest (6) and Wooded Sandbank Vegetation (49). A distinct F+S grouping additionally includes Savanna Formation (4) and Flooded Savanna (7). Plantations (9) remain separate. Published crop subclasses, pasture and mosaics are retained as distinct destinations. [2]

Indonesia uses the complete Collection 4.1 workbook: 441 leaf-class province rows, 38 provinces, seven island groups and an approximately 189.044-million-hectare reported footprint. Its official mapped Forest group comprises Forest Formation (3), Mangrove (5) and Peat Swamp Forest (76). Oil palm (35) and pulpwood plantations (9) are agricultural classes, not members of this forest group. The comparison also examines the earlier Collection 3 table, excluding its separately labelled 20-kilometre island-buffer category. Collection 3 and 4.1 have different reported geographic support, so their contrast is a collection-and-support sensitivity, not a clean version effect. [3,4]

Leaf classes are summed only once. Annual transition row and column margins are checked against the corresponding coverage areas at the finest reported geographic/class level. Forest inflow minus outflow is reconciled with each annual extent change. Maximum margin residuals are below 0.000005 ha in Brazil and 0.0000003 ha in Indonesia; these are numerical consistency checks, not accuracy estimates. Complete provider partitions do not certify exact equivalence to every national-account component's territorial convention. Zero tabulated no-data area also does not establish cloud-free observation: gap filling and temporal filtering can assign complete labels.

### 2.3. Unique endpoint conversions from annual aggregate capacities

Annual outflows count transition events, including land that changes more than once. They are not unique 2015-to-2023 forest conversion hectares. To learn more without inventing pixel trajectories, we construct a layered network with class nodes for each year and observed annual class-to-class area T_(t,i,j) as edge capacity. A marked flow x_(t,i,j) represents land mapped as forest in 2015. On the first transition x = T for forest-origin edges and x = 0 otherwise. At every interior class-year node, marked inflow equals marked outflow, with 0 ≤ x ≤ T. Minimizing and maximizing marked inflow into a specified 2023 target class gives its feasible endpoint-conversion range.

These bounds are sharp for the ideal balanced, continuous-area aggregate-capacity model: every feasible marked flow, and its unmarked complement, can be decomposed into paths on the acyclic network. Conversely, any pixel history consistent with the aggregate matrices induces such a feasible flow. The calculation therefore bounds endpoint cohort areas without claiming to recover the actual histories. Source rounding and numerical tolerances are recorded separately. Solving within fixed reported geographic units and summing their extrema restricts cross-region rearrangements and cannot weaken the raw bounds. We evaluate national aggregation and 49 Brazilian state-biome intersections or 38 Indonesian provinces, yielding 634 linear programmes. Saved primal and dual witnesses are independently checked from the deposited capacities, without re-solving.

The target is land mapped as F in 2015 and, for example, oil palm in 2023; it can pass through intermediate classes and can experience repeated conversion. This is not necessarily a single direct forest-to-oil-palm event. It is not first clearing, a calibrated true transition, converted biomass or a causal effect.

### 2.4. Error requirements and covariance-independent inventory inference

If each mapped endpoint area has absolute area error bounded by e, the true area difference lies within mapped ΔA ± 2e. A sign ceases to be guaranteed at e = |ΔA|/2. These break-even requirements are conditional sensitivity results, not measured uncertainty intervals. For an endpoint-conversion intersection, the appropriate errors are instead symmetric differences of the initial forest and final destination class sets. If their area-error budgets sum to E, a mapped lower conversion bound L implies a reference-defined lower bound no smaller than max(0,L−E). A bound on net class area alone does not supply this class-set error budget.

Finnish NFI12 (2014–2018) and NFI13 (2019–2023) refer to productive and poorly productive forest across the whole country, including Åland. The national APIs supply biomass components, volume and area; the original inventory supplementary tables supply volume standard errors. These inventory cycles are not relabelled exact 2015 and 2023 observations. [6,7]

For two marginal approximately normal inventory estimators with standard errors s_0 and s_1, form two-sided 97.5% endpoint intervals. Bonferroni's union bound gives at least approximately 95% simultaneous endpoint coverage without assuming independence or a particular joint distribution. Subtracting the endpoints gives ΔV ± z_0.9875(s_0+s_1), where z_0.9875 = 2.241403. A small outward allowance covers published point/SE rounding. This addresses sampling uncertainty under marginal normal approximation, not unbounded differential measurement or model bias. Separately, the sensitivity ΔV ± √20(s_0+s_1) uses a worst-covariance second-moment bound, conditional on unbiased errors and treating the reported standard errors as true error standard deviations. Its stronger distribution-free interpretation is not asserted for estimated variances.

## 3. Results

### 3.1. Actual account trajectories and domain dependence

Total account biocapacity increases in 91 of the 183 eligible identifiers and decreases in 92. Of the increasing totals, 68 coexist with a declining forest account component, 22 with an increasing component and one with a component change unresolved at the rounding scale. This is within-account composition, not a count of ecological asset losses. All eight inherited candidates are retained in Table 1.

Table 1. Actual 2015–2023 percentage changes. “− both” excludes fishing and built components; forest is the account component, not mapped area or physical biomass. Per-person results are separate quantities.

{account_table}

Cameroon's total-account sign changes when built area is excluded; Canada's changes when fishing is excluded. Indonesia remains increasing under all three exclusions, although its per-person total declines. Brazil's total decreases, so it does not exemplify an increasing aggregate account concealing loss. The candidate set is not reduced to the cases with the desired direction.

![Figure 1. Account composition for all eight candidates. Both series are accounting quantities; neither is an independent measure of ecological stock.](figures/Figure1_account_composition_F12.png)

### 3.2. National mapped forest extent

Brazil's formations decline from 412.016 to 401.154 million hectares: −10.863 million hectares, or −2.64%. The broader formations-plus-savanna group declines by 18.184 million hectares (−3.43%). Brazil's total biocapacity declines by 1.21%. The two negative signs do not establish equality of ecological condition, but they contradict an increasing-total narrative for this country and window.

Indonesia's Collection 4.1 mapped forest declines from 107.055 to 106.352 million hectares: −702,562 ha, or −0.66%, while total biocapacity increases by 5.46%. Excluding fishing gives +7.57%; excluding fishing and built gives +6.68%. This is a nominal national account–mapped-extent discordance, not yet a calibrated national biomass loss. The combined forest result is definition-sensitive: Forest Formation alone increases by 3,019 ha (+0.0032%), mangrove decreases by 11,620 ha, and peat swamp forest decreases by 693,961 ha (−6.35%). Class components describe the classification, not identified physical causes.

Collection 3 reports a larger Indonesian forest decline of 3.106 million hectares (−3.17%) over the same endpoint years. Its retained all-class footprint exceeds the Collection 4.1 reported footprint by 43,130 ha, but that difference in total area is not a bound on their geometric symmetric difference. The 2.403-million-hectare difference between reported net forest changes cannot therefore be attributed solely to version processing. Both collections give a negative combined forest change, but they differ materially in levels and magnitude.

![Figure 2. Indonesia: operational forest definitions and collection-and-support sensitivity. These contrasts are not a confidence interval or a pure map-version effect.](figures/Figure2_Indonesia_definition_source_F12.png)

### 3.3. Gross mapped events, spatial compensation and endpoint bounds

Brazilian formations have 22.711 million hectare-events of annual outflow and 11.848 million hectare-events of annual inflow. Their difference matches the −10.863-million-hectare endpoint change. Indonesian forest has 4.424 million hectare-events of outflow and 3.722 million hectare-events of inflow, matching the −0.703-million-hectare net change. Summing repeated-year event areas is not a unique land-area estimate.

Across Indonesia's provinces, the sum of negative province-level net changes is {float(idr['sum_negative_unit_net_changes_ha'])/1e6:.3f} million hectares and the sum of positive province-level net changes is {float(idr['sum_positive_unit_net_changes_ha'])/1e6:.3f} million hectares. Spatially aggregated net change therefore omits substantial classified regional compensation. These are descriptive map quantities, not confidence-resolved regional losses and gains.

Table 2. Conditional unique 2015-to-2023 mapped endpoint-cohort ranges using the finest reported geographic constraints. F denotes each country's stated forest/formation group; F+S denotes Brazilian formations plus savannas. Displayed endpoints are rounded; full-precision data include outward numerical allowances. These are feasible ranges, not confidence intervals.

{path_table}

Indonesia's mapped 2015-forest-to-2023-oil-palm area is bounded below by about 111,612 ha and above by about 2.326 million ha. The lower bound remains positive in the classified aggregate model despite allowing every compatible intervening path. The range is not the 707,004 hectare-events of direct annual forest-to-oil-palm transitions: endpoint cohorts can pass through other classes. Brazil's formations-to-cropland lower bound is zero, whereas the formations-plus-savanna-to-cropland lower bound is about 9,386 ha. Thus the annual tables support some geographically constrained endpoint statements but do not identify all proposed conversion mechanisms or remove dependence on the forest definition.

![Figure 3. Geographically refined mapped endpoint-conversion ranges. A positive mapped lower bound is not automatically a positive reference-defined ecological conversion bound.](figures/Figure3_endpoint_conversion_bounds_F12.png)

### 3.4. Finland: a supported secondary volume result

The national inventory API reports total living-tree biomass of 1,744 and 1,786 million tonnes for NFI12 and NFI13, respectively. This includes stumps and roots. The sum of published stems plus branches/foliage is 1,350 and 1,384 million tonnes; independently rounded components can differ from the published total by approximately one million tonnes. Neither series is silently equated with the primary remote-sensing above-ground biomass definition, and the API table supplies no marginal biomass standard errors.

The original volume/SE cells give 2,475.192 ± 13.001 million m³ for NFI12 and 2,551.964 ± 13.534 million m³ for NFI13, where ± denotes one reported sampling standard error, not a 95% interval. The difference is +76.772 million m³ (+3.10%). The approximate 95% simultaneous-endpoint Bonferroni difference interval is [{lo:.1f}, {hi:.1f}] million m³, including the conservative source-rounding allowance. Cross-cycle independence is not assumed. A differential-bias budget of approximately {lo:.1f} million m³ would exhaust the positive lower sampling bound; that bias budget has not been independently measured.

The second-moment sensitivity interval is [{mlo:.1f}, {mhi:.1f}] million m³ and does not resolve the sign. The positive conclusion is consequently an approximate design-based sampling result for the two inventory cycles, not an assumption-free statement about total error or an exact 2015–2023 biomass result. Geographic inclusion of Åland is established for these inventory tables, while equivalence to every account component's territorial convention remains a separate question.

![Figure 4. Finnish inventory-cycle volume inference. The approximate marginal-normal interval resolves an increase without assuming independence; the stronger moment-only sensitivity does not. Neither quantifies all systematic error.](figures/Figure4_Finland_sampling_assumptions_F12.png)

Official forest-account timber stocks provide a further warning about temporal and accounting operators. In 2023, opening-to-closing reported volume falls by 4.185 million m³ even though net increment minus removals is +20.802 million m³, because the balancing item is −24.987 million m³. Reclassification is zero and the reconciliation closes. This shorter-window, partly inventory-derived forest account is neither an independent biomass validation nor proof of observed ecological loss from the negative closing-stock difference. [8]

## 4. Validation limits are quantitative requirements, not blanket exclusions

For Brazil's formations, symmetric endpoint absolute-area errors must each be smaller than 5.431 million ha to guarantee the negative net sign under that bound; this is 0.6385% of the tabulated national footprint per endpoint. For Indonesia's combined forest, the corresponding requirement is 351,281 ha, or 0.1858% of its footprint per endpoint. No measured endpoint error budget satisfying either requirement is claimed.

The Indonesian forest-to-oil-palm lower endpoint bound requires a different calibration: the combined initial-forest and final-oil-palm class-set symmetric-difference error budget must be below about 111,612 ha to retain a positive reference-defined lower bound. Equal endpoint allocations would be about 55,806 ha each, or 0.0295% of the reported footprint. A pooled or incompatible overall accuracy percentage, or even a small error in net class area, does not provide that set-error guarantee.

The current Brazilian Collection 11 accuracy appendix describes reference collection 3 for 1985–2024, with boundary-point exclusions, interpreter-agreement weighting, sparse-class exclusions and map-dependent legend harmonization. The public collection-2.1 point representation examined here is not that same reference vintage. Of its 85,152 points, {refs['paired_nonblank']:,} have nonblank 2015 and 2023 labels, {refs['paired_observed_nontraining']:,} remain after excluding non-observed labels and the archive's training role, and {refs['paired_observed_nontraining_nonedge_both']:,} are also nonedge at both endpoints. Their geographic distribution is highly uneven; nonblank labels alone do not establish national representativeness or compatibility. The old archive's 9,999 training tokens also differ from the newer document's stated 10,000 training points, which is retained as a provenance discrepancy rather than silently reconciled. [5]

Indonesia's current Collection 4.1 ATBD lists no global accuracy estimate for that collection in its collection summary. The described independent annual reference labels cover 2000–2022, and the public accuracy page describes Collection 2.0. These facts do not show that newer validation is impossible or nonexistent; they show why the acquired documentation and earlier pooled accuracies cannot certify the 2015–2023 Collection 4.1 change. Marginal reference-based endpoint intervals can suffice: exact paired censuses and known covariance are not mandatory if compatible designs yield adequate simultaneous bounds. [4]

## 5. Implications and remaining empirical questions

The results establish actual national account changes, national published-map extent dynamics and conditional conversion ranges for two countries, plus a defensible secondary inventory-cycle volume result. They no longer support treating all national evidence as absent. Equally, they do not establish any of the eight original calibrated national 2015–2023 above-ground biomass changes. The Democratic Republic of the Congo, Cameroon, Canada, France and Australia have not received completed physical national analyses in this extension; that is unfinished work, not a demonstrated permanent lack of suitable evidence.

No conversion association, causal conversion effect, or reporting false-negative rate against independently supported ecological outcomes is estimated. Repeating a single national account change as if it supplied independent state-level observations would not solve that problem. The two processed map cases are also not a representative country sample. The nominal Indonesian discordance and Brazilian concordance are retained together, without promoting their one-of-two pattern into population accuracy.

A useful reporting sidecar would present the aggregate account with its components and domain sensitivities; a separately defined physical extent or stock; a source-appropriate uncertainty or feasible set; and conversion diagnostics that distinguish repeated events from unique endpoint cohorts. Such a sidecar supplies information but does not by itself demonstrate improved reporting accuracy. That claim requires independent reference outcomes on common cases. The useful distinction is not a choice between accounts and assets, but a requirement to show both and identify the evidential status of each.

## Data, methods and assistance

Machine-readable results, source URLs and SHA-256 identities, original figures, analysis code, compressed annual transition capacities and saved primal/dual witnesses accompany this extension. Provider workbook and PDF bodies are not redistributed. The workbook national_results_F12_A01.xlsx supplies directly readable result tables and definitions. The separate completion account documents executed dependencies and unfinished work; it is not evidence of ecological validation.

Arena.ai Agent Mode was used to assist with computational methods and literature review.

## References

[1] Footprint Data Foundation and Global Footprint Network. National Ecological Footprint and Biocapacity Accounts, 2026 edition, results and metadata v1.0. Exact workbook identity is supplied in the source metadata. https://data.footprintnetwork.org/

[2] MapBiomas Brazil. Collection 11 national coverage and transition statistics, biome-state workbook, accessed 30 October 2026. https://brasil.mapbiomas.org/iniciativas-e-produtos/cobertura-e-uso-da-terra/cobertura-30m/cobertura/

[3] MapBiomas Indonesia. Collection 4.1 and Collection 3 land-cover statistics, accessed 30 October 2026. https://landy.mapbiomas.id/en/statistics

[4] MapBiomas Indonesia. Algorithm Theoretical Basis Document, Collection 4.1, February 2026, especially pp. 5, 12 and 28–29; and public accuracy-assessment page. https://landy.mapbiomas.id/en/atbd and https://landy.mapbiomas.id/en/accuracy-assessment

[5] MapBiomas Brazil. Accuracy Assessment Appendix, Collection 11 v1, sections 2.2.3–2.2.4; public reference-point downloads. https://brasil.mapbiomas.org/downloads/pontos-de-validacao/

[6] Korhonen, K.T., Ahola, A., et al. (2021). Forests of Finland 2014–2018 and their development 1921–2018. Silva Fennica 55, article 10662. Supplementary S2, Table 20a. https://doi.org/10.14214/sf.10662

[7] Korhonen, K.T., et al. (2024). Forests of Finland 2019–2023 and their development 1921–2023. Silva Fennica 58, article 24045. Supplementary S2, Table 55. https://doi.org/10.14214/sf.24045

[8] Natural Resources Institute Finland (Luke). Archived forest resources statistics and forest accounts, national inventory biomass/volume/area and 2022–2023 timber-stock tables. API queries and returned data are deposited. https://www.luke.fi/en/statistics/forest-resources
'''
 (ROOT/'national_empirical_extension_F12_A01.md').write_text(text);return text


def pdf(text):
 from matplotlib.font_manager import findfont,FontProperties
 for name,weight in [('Body','normal'),('Bold','bold')]:pdfmetrics.registerFont(TTFont(name,findfont(FontProperties(family='DejaVu Sans',weight=weight))))
 styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='B',fontName='Body',fontSize=9.3,leading=13.4,textColor=colors.HexColor(NAVY),spaceAfter=7));styles.add(ParagraphStyle(name='T',fontName='Bold',fontSize=21,leading=26,textColor=colors.HexColor(NAVY),spaceAfter=14));styles.add(ParagraphStyle(name='H',fontName='Bold',fontSize=13,leading=17,textColor=colors.HexColor(NAVY),spaceBefore=13,spaceAfter=8,keepWithNext=True));styles.add(ParagraphStyle(name='S',fontName='Bold',fontSize=10.4,leading=14,textColor=colors.HexColor(TEAL),spaceBefore=10,spaceAfter=6,keepWithNext=True));styles.add(ParagraphStyle(name='C',fontName='Body',fontSize=8,leading=11,textColor=colors.HexColor(GREY),spaceAfter=11));styles.add(ParagraphStyle(name='TC',fontName='Body',fontSize=8,leading=11,textColor=colors.HexColor(GREY),spaceAfter=7,keepWithNext=True));styles.add(ParagraphStyle(name='R',fontName='Body',fontSize=8.2,leading=11.3,textColor=colors.HexColor(NAVY),spaceAfter=6));styles.add(ParagraphStyle(name='Cell',fontName='Body',fontSize=7.4,leading=10));story=[];lines=text.splitlines();i=0
 def para(s,style='B'):
  s=escape(s);s=re.sub(r'(https://[^\s]+)',lambda m:'<link href="'+m.group(0)+'" color="#16847b">'+m.group(0)+'</link>',s);return Paragraph(s,styles[style])
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('!['):
   m=re.match(r'!\[(.*?)\]\((.*?)\)',line);from PIL import Image as PILImage
   p=ROOT/m.group(2)
   with PILImage.open(p)as im:w,h=im.size
   story.append(KeepTogether([Image(str(p),width=502,height=502*h/w),Spacer(1,4),para(m.group(1),'C')]));i+=1;continue
  if line.startswith('|'):
   block=[]
   while i<len(lines)and lines[i].strip().startswith('|'):block.append([v.strip()for v in lines[i].strip().strip('|').split('|')]);i+=1
   block=[r for r in block if not all(set(x)<=set('-: ')for x in r)];n=len(block[0]);widths=[502/n]*n
   if n==4:widths=[161,151,95,95]
   cells=[[para(v,'Cell')for v in r]for r in block];t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7efef')),('LINEBELOW',(0,0),(-1,0),.7,colors.HexColor(TEAL)),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f6f8')]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));caption=story.pop()if story and isinstance(story[-1],Paragraph)and story[-1].style.name=='TC'else None;story.append(KeepTogether(([caption]if caption else[])+[t,Spacer(1,10)]));continue
  if line.startswith('# '):story.append(para(line[2:],'T'));i+=1;continue
  if line.startswith('### '):story.append(para(line[4:],'S'));i+=1;continue
  if line.startswith('## '):story.append(para(line[3:],'H'));i+=1;continue
  block=[line];i+=1
  while i<len(lines)and lines[i].strip()and not lines[i].startswith(('#','|','![')):block.append(lines[i].strip());i+=1
  story.append(para(' '.join(block),'TC'if line.startswith('Table ')else'C'if line.startswith('Figure ')else'R'if line.startswith('[')else'B'))
 def footer(c,doc):
  c.setStrokeColor(colors.HexColor('#d6dfe5'));c.line(46,36,548,36);c.setFont('Body',8);c.setFillColor(colors.HexColor(GREY));c.drawString(46,24,'NATIONAL EMPIRICAL EXTENSION · F12_A01');c.drawRightString(548,24,str(doc.page))
 doc=SimpleDocTemplate(str(ROOT/'national_empirical_extension_F12_A01.pdf'),pagesize=A4,rightMargin=47,leftMargin=46,topMargin=44,bottomMargin=48,title='National account composition, mapped forest extent and inventory-cycle stock change',author='Amin Abaee',pageCompression=1);doc.build(story,onFirstPage=footer,onLaterPages=footer)


def excel():
 wb=Workbook();about=wb.active;about.title='Read_me'
 content=[['F12 national evidence results','30 October 2026'],['Quantities are distinct','gha = accounting; ha = mapped area; Mt = dry biomass; Mm3 = growing-stock volume.'],['Primary target','2015–2023 national above-ground biomass change remains unadmitted for all eight candidates.'],['Completed national map results','BRA C11 and IDN C4.1 extent, event balances and conditional endpoint-conversion bounds.'],['Finnish inference','Secondary inventory-cycle growing-stock volume; approximate marginal-normal Bonferroni interval, not exact-endpoint AGB.'],['Conversion intervals','Feasible mapped ranges, not confidence intervals or calibrated true conversions.'],['Error budgets','Break-even requirements are NOT measured uncertainty. Class-set error is not net-area error.'],['Source sensitivity','Indonesia C3 versus C4.1 mixes collection and reported-support differences.'],['Account sensitivities','Primary 183 identifiers; secondary 2010–2020 has 182, SSD unavailable.'],['Selection','Exploratory feasibility-led analysis; no independent held-out or representative-country claim.'],['Author ORCID','https://orcid.org/0000-0002-0019-1842'],['AI declaration','Arena.ai Agent Mode was used to assist with computational methods and literature review.']]
 for r in content:about.append(r)
 sheets=[('Accounts_8','accounts_candidate_changes_F12.csv'),('Accounts_183','accounts_changes_183_F12.csv'),('National_cover','national_account_mapped_cover_comparisons_F12.csv'),('Extent_BRA','BRA_C11_extent_changes_F12.csv'),('Extent_IDN','IDN_C41_extent_changes_F12.csv'),('Unique_conversions','endpoint_conversion_path_bounds_F12.csv'),('Conversion_error','endpoint_conversion_set_error_requirements_F12.csv'),('Event_sums','national_mapped_transition_event_sums_F12.csv'),('Regional_compensation','national_mapped_regional_compensation_F12.csv'),('FIN_volume_SE','FIN_inventory_volume_SE_source_cells_F12.csv'),('FIN_biomass','FIN_inventory_biomass_operational_summaries_F12.csv'),('FIN_bias_sensitivity','FIN_cycle_volume_bias_sensitivity_F12.csv'),('FIN_stock_flow','FIN_forest_account_stock_flow_balance_F12.csv'),('IDN_C3_sensitivity','IDN_C3_national_extent_sensitivity_F12.csv'),('All_eight_status','national_evidence_status_all_eight_F12.csv'),('Reference_completeness','BRA_reference_v21_label_completeness_F12.csv')]
 def typed(v):
  if v=='':return None
  if v in ['True','False']:return v=='True'
  try:return float(v)if any(x in v.lower()for x in ['.','e'])else int(v)
  except ValueError:return v
 for name,file in sheets:
  r=rows(file);ws=wb.create_sheet(name);ws.append(list(r[0]));
  for row in r:ws.append([typed(v)for v in row.values()])
 j=load('FIN_cycle_volume_inference_F12.json');ws=wb.create_sheet('FIN_inference')
 ws.append(['Quantity','Value / explicit condition'])
 for k,v in j.items():ws.append([k,json.dumps(v)if isinstance(v,(list,dict))else v])
 ws=wb.create_sheet('Sources');ws.append(['File / source','URL','SHA256','Bytes','Raw redistributed'])
 for name,r in load('source_registry_F12.json').items():ws.append([name,r['url'],r['sha256'],r['bytes'],r['raw_redistributed']])
 for name in ['FIN_inventory_biomass','FIN_inventory_volume','FIN_inventory_area','FIN_forest_account_volume']:
  q=load(name+'_query_F12.json');ws.append([name,q['url'],q['response_sha256'],None,'Public statistical response supplied; workbook/PDF bodies not supplied'])
 for ws in wb:
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False;ws.row_dimensions[1].height=33
  for c in ws[1]:c.font=Font(name='Aptos',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17344A');c.alignment=Alignment(wrap_text=True,vertical='center')
  for col in ws.columns:
   letter=col[0].column_letter;ws.column_dimensions[letter].width=42 if ws.title in ['Read_me','FIN_inference']else min(34,max(14,len(str(col[0].value))*.75))
   for c in col[1:]:
    c.font=Font(name='Aptos',size=10);c.alignment=Alignment(vertical='top',wrap_text=isinstance(c.value,str)and len(c.value)>40)
    if isinstance(c.value,float):c.number_format='#,##0.0000;[Red]-#,##0.0000'
    if isinstance(c.value,str)and c.value.startswith('https://'):c.hyperlink=c.value;c.style='Hyperlink'
  if ws.title in ['Read_me','FIN_inference']:ws.column_dimensions['A'].width=49;ws.column_dimensions['B'].width=105
 wb.save(ROOT/'national_results_F12_A01.xlsx')

if __name__=='__main__':
 figures();text=manuscript();pdf(text);excel();print('Scientific extension, four figures and Excel results workbook created.')
