#!/usr/bin/env python3
"""PDF text/geometry/link checks, full-page rendering and Excel structure/value QA."""
import csv,hashlib,json,math
from pathlib import Path
import pymupdf
from PIL import Image,ImageOps,ImageDraw
from openpyxl import load_workbook
from accounts_F12 import require
ROOT=Path(__file__).resolve().parents[1];CACHE=ROOT.parent/'.cache/F12/visual';CACHE.mkdir(parents=True,exist_ok=True)


def run(confirmed=False):
 p=ROOT/'national_empirical_extension_F12_A01.pdf';d=pymupdf.open(p);texts=[];outside=[];links=[];tiles=[]
 for i,page in enumerate(d):
  texts.append(page.get_text());links.extend(x.get('uri','')for x in page.get_links())
  for b in page.get_text('blocks'):
   x0,y0,x1,y1=b[:4]
   if x0<-.2 or y0<-.2 or x1>page.rect.width+.2 or y1>page.rect.height+.2:outside.append(dict(page=i+1,bbox=b[:4]))
  pix=page.get_pixmap(matrix=pymupdf.Matrix(.45,.45));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);canvas=Image.new('RGB',(im.width+12,im.height+28),'#dfe6eb');canvas.paste(im,(6,6));ImageDraw.Draw(canvas).text((8,im.height+10),str(i+1),fill='black');tiles.append(canvas)
 text='\n'.join(texts);require(not outside,'PDF text outside page');require('�'not in text,'PDF replacement glyph');require('0000-0002-0019-1842'in text and any('orcid.org/0000-0002-0019-1842'in x for x in links),'ORCID hyperlink absent');require('Arena.ai Agent Mode was used to assist with computational methods and literature review.'in text.replace('\n',' '),'Exact declaration absent');require(all(len(t.strip())>100 for t in texts),'Near-empty page')
 cols=4;sheet=Image.new('RGB',(cols*tiles[0].width,math.ceil(len(tiles)/cols)*tiles[0].height),'#cdd7df')
 for i,im in enumerate(tiles):sheet.paste(im,((i%cols)*im.width,(i//cols)*im.height))
 sheet.save(CACHE/'report_all_pages_F12.png')
 workbook=ROOT/'national_results_F12_A01.xlsx';wb=load_workbook(workbook,data_only=False);require(len(wb.sheetnames)==19,'Unexpected workbook sheets');require(all(ws.freeze_panes=='A2'for ws in wb),'Missing frozen headers');require(not any(c.data_type=='f'for ws in wb for row in ws for c in row),'Unexpected spreadsheet formula')
 checks=[('Accounts_8','accounts_candidate_changes_F12.csv'),('Accounts_183','accounts_changes_183_F12.csv'),('Unique_conversions','endpoint_conversion_path_bounds_F12.csv'),('FIN_volume_SE','FIN_inventory_volume_SE_source_cells_F12.csv')];cells=0
 for sheetname,filename in checks:
  with (ROOT/'data'/filename).open()as f:r=list(csv.reader(f))
  ws=wb[sheetname];require(ws.max_row==len(r),'Spreadsheet row count mismatch')
  for y,row in enumerate(r,1):
   for x,v in enumerate(row,1):
    actual=ws.cell(y,x).value
    if v=='':require(actual is None,'Expected blank spreadsheet cell')
    elif v in ['True','False']:require(actual==(v=='True'),'Boolean mismatch')
    else:
     try:expected=float(v)
     except ValueError:require(str(actual)==v,'Text mismatch in spreadsheet')
     else:require(isinstance(actual,(int,float))and math.isclose(actual,expected,rel_tol=2e-13,abs_tol=1e-8),'Numeric mismatch in spreadsheet')
    cells+=1
 result=dict(pdf=dict(file=p.name,pages=len(d),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),text_outside_page_bounds=outside,replacement_characters=False,exact_AI_declaration=True,ORCID_hyperlink=True,all_pages_rendered=True,contact_sheet='.cache/F12/visual/report_all_pages_F12.png'),workbook=dict(file=workbook.name,sheets=len(wb.sheetnames),sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),compared_cells=cells,unexpected_formulas=False),visual_review_completed=confirmed,visual_review_pages=list(range(1,len(d)+1))if confirmed else[])
 (ROOT/'metadata/deliverable_QA_F12.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))

if __name__=='__main__':
 import argparse
 p=argparse.ArgumentParser();p.add_argument('--confirm-visual-review',action='store_true');run(p.parse_args().confirm_visual_review)
