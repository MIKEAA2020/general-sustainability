#!/usr/bin/env python3
"""Rebuild ECOMOD v45's NFA 2026 world-series tables and Figures 2--3.

Primary window: complete published world observations, 1961--2023.
Sensitivity: 1961--2022 from the SAME NFA 2026 edition and population series.
The 2024--2025 workbook rows are estimates and are excluded.

Source workbook: data/NEFBA_2026.xlsx, Edition 2026, results/metadata v1.0.
Population: world_data.pop_1000_persons (UN PopStat), converted to persons.
"efp_total_gha" is the global Ecological Footprint of Production; it is not
silently relabeled as the consumption footprint. NFA land-type values are
account values, not observations of physical ecological stocks.

Run from any directory: python work/nfa_2026_recompute.py
Requires openpyxl, numpy, and matplotlib.
"""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "NEFBA_2026.xlsx"
DATA_DIR = ROOT / "data"
FIG2 = ROOT / "fig2_nfa_aggregation_face.png"
FIG3 = ROOT / "fig3_composition_attribution.png"

EDITION = "National Ecological Footprint and Biocapacity Accounts, Edition 2026, results and metadata v1.0"
POPULATION_SOURCE = "world_data.pop_1000_persons; UN PopStat; converted from 1,000 persons to persons"
COMPONENTS = {
    "Cropland": "biocap_crop_gha",
    "Grazing land": "biocap_grazing_gha",
    "Forest biocapacity": "biocap_forest_gha",
    "Fishing grounds": "biocap_fishing_gha",
    "Built-up land": "biocap_built_gha",
}
AREA_COLUMNS = {
    "area_crop_ha": "area_crop_ha",
    "area_grazing_ha": "area_grazing_ha",
    "area_forest_ha": "area_forest_ha",
    "area_fishing_ha": "area_fishing_ha",
    "area_built_ha": "area_built_ha",
    "area_total_ha": "area_total_ha",
}


def logmean(a: float, b: float) -> float:
    """Logarithmic mean of positive endpoint values."""
    if a <= 0 or b <= 0:
        raise ValueError("Log-mean requires positive endpoints")
    if math.isclose(a, b, rel_tol=1e-13, abs_tol=0.0):
        return (a + b) / 2
    return (b - a) / math.log(b / a)


def load_world() -> list[dict]:
    if not WORKBOOK.is_file():
        raise FileNotFoundError(f"Missing official source workbook: {WORKBOOK}")
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if "world_data" not in wb.sheetnames:
        raise ValueError(f"Unexpected workbook sheets: {wb.sheetnames}")
    ws = wb["world_data"]
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if row.get("country_name") == "World" and row.get("year") is not None:
            rows.append(row)
    years = [int(r["year"]) for r in rows]
    if len(years) != len(set(years)):
        raise ValueError("Duplicate world-year rows in NFA workbook")
    if not set(range(1961, 2024)).issubset(set(years)):
        raise ValueError("The 2026 workbook lacks the expected complete 1961--2023 world series")
    by_year = {int(r["year"]): r for r in rows}
    # 2023 is an observed/completed row; 2024--2025 are estimates, not the primary window.
    if by_year[2023]["flag"] == "E":
        raise ValueError("2023 is flagged estimated; check the 2026 edition description")
    for year in (2024, 2025):
        if year in by_year and by_year[year]["flag"] != "E":
            raise ValueError(f"Expected {year} to be flagged estimated in the 2026 workbook")
    selected = [by_year[y] for y in range(1961, 2024)]
    required = {
        "pop_1000_persons", "efp_total_gha", "biocap_total_gha",
        *COMPONENTS.values(), *AREA_COLUMNS.values(),
    }
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Missing required 2026 NFA columns: {sorted(missing)}")
    for row in selected:
        if row["pop_1000_persons"] is None or float(row["pop_1000_persons"]) <= 0:
            raise ValueError(f"Invalid population value for {row['year']}")
    return selected


def enrich(row: dict) -> dict:
    out = {
        "year": int(row["year"]),
        "quality_flag": row["flag"],
        "population_persons": float(row["pop_1000_persons"]) * 1000.0,
        "population_1000_persons": float(row["pop_1000_persons"]),
        "E_production_gha": float(row["efp_total_gha"]),
        "B_gha": float(row["biocap_total_gha"]),
    }
    p = out["population_persons"]
    out["E_production_percap_gha_person"] = out["E_production_gha"] / p
    out["B_percap_gha_person"] = out["B_gha"] / p
    out["E_production_over_B"] = out["E_production_gha"] / out["B_gha"]
    for label, col in COMPONENTS.items():
        out[f"B_{label.lower().replace(' ', '_')}_gha"] = float(row[col])
    for label, col in AREA_COLUMNS.items():
        out[label] = float(row[col])
    return out


def endpoint_metrics(rows: list[dict], end_year: int) -> dict:
    by_year = {r["year"]: r for r in rows}
    start_year = 1961
    a, b = by_year[start_year], by_year[end_year]
    dln_b = math.log(b["B_gha"] / a["B_gha"])
    dln_p = math.log(b["population_persons"] / a["population_persons"])
    dln_bpc = math.log(b["B_percap_gha_person"] / a["B_percap_gha_person"])
    dln_epc = math.log(b["E_production_percap_gha_person"] / a["E_production_percap_gha_person"])
    first_overshoot = next(
        r["year"] for r in rows
        if start_year <= r["year"] <= end_year and r["E_production_over_B"] > 1
    )
    component_sum_start = sum(a[f"B_{k.lower().replace(' ', '_')}_gha"] for k in COMPONENTS)
    component_sum_end = sum(b[f"B_{k.lower().replace(' ', '_')}_gha"] for k in COMPONENTS)
    return {
        "period": f"{start_year}-{end_year}",
        "start_year": start_year,
        "end_year": end_year,
        "edition": EDITION,
        "population_source": POPULATION_SOURCE,
        "B_start_gha": a["B_gha"],
        "B_end_gha": b["B_gha"],
        "B_endpoint_ratio": b["B_gha"] / a["B_gha"],
        "B_change_pct": 100 * (b["B_gha"] / a["B_gha"] - 1),
        "dln_B": dln_b,
        "population_start_persons": a["population_persons"],
        "population_end_persons": b["population_persons"],
        "population_endpoint_ratio": b["population_persons"] / a["population_persons"],
        "population_change_pct": 100 * (b["population_persons"] / a["population_persons"] - 1),
        "dln_population": dln_p,
        "B_percap_start_gha_person": a["B_percap_gha_person"],
        "B_percap_end_gha_person": b["B_percap_gha_person"],
        "B_percap_endpoint_ratio": b["B_percap_gha_person"] / a["B_percap_gha_person"],
        "B_percap_change_pct": 100 * (b["B_percap_gha_person"] / a["B_percap_gha_person"] - 1),
        "dln_B_percap": dln_bpc,
        "dln_identity_residual": dln_b - dln_p - dln_bpc,
        "E_production_percap_start_gha_person": a["E_production_percap_gha_person"],
        "E_production_percap_end_gha_person": b["E_production_percap_gha_person"],
        "E_production_percap_endpoint_ratio": b["E_production_percap_gha_person"] / a["E_production_percap_gha_person"],
        "E_production_percap_change_pct": 100 * (b["E_production_percap_gha_person"] / a["E_production_percap_gha_person"] - 1),
        "dln_E_production_percap": dln_epc,
        "E_production_over_B_start": a["E_production_over_B"],
        "E_production_over_B_end": b["E_production_over_B"],
        "first_year_E_production_over_B_gt_1": first_overshoot,
        "component_sum_residual_start_gha": component_sum_start - a["B_gha"],
        "component_sum_residual_end_gha": component_sum_end - b["B_gha"],
    }


def component_metrics(rows: list[dict], end_year: int) -> list[dict]:
    by_year = {r["year"]: r for r in rows}
    a, b = by_year[1961], by_year[end_year]
    out = []
    for label in COMPONENTS:
        key = f"B_{label.lower().replace(' ', '_')}_gha"
        start, end = a[key], b[key]
        out.append({
            "period": f"1961-{end_year}",
            "start_year": 1961,
            "end_year": end_year,
            "component": label,
            "account_value_start_gha": start,
            "account_value_end_gha": end,
            "change_gha": end - start,
            "endpoint_ratio": end / start,
            "contribution_pp_of_B1961": 100 * (end - start) / a["B_gha"],
            "unit_and_status": "gha; annual NFA biocapacity account value, not a physical-stock measurement",
        })
    return out


def aggregate_component_metrics(rows: list[dict], end_year: int) -> dict:
    by_year = {r["year"]: r for r in rows}
    a, b = by_year[1961], by_year[end_year]
    start_crop = a["B_cropland_gha"]
    end_crop = b["B_cropland_gha"]
    start_non_crop = a["B_gha"] - start_crop
    end_non_crop = b["B_gha"] - end_crop
    return {
        "period": f"1961-{end_year}",
        "start_year": 1961,
        "end_year": end_year,
        "crop_biocapacity_start_gha": start_crop,
        "crop_biocapacity_end_gha": end_crop,
        "crop_biocapacity_endpoint_ratio": end_crop / start_crop,
        "crop_biocapacity_change_pct_of_crop_start": 100 * (end_crop / start_crop - 1),
        "crop_area_start_ha": a["area_crop_ha"],
        "crop_area_end_ha": b["area_crop_ha"],
        "crop_area_endpoint_ratio": b["area_crop_ha"] / a["area_crop_ha"],
        "NFA_crop_account_value_per_ha_start_gha_ha": start_crop / a["area_crop_ha"],
        "NFA_crop_account_value_per_ha_end_gha_ha": end_crop / b["area_crop_ha"],
        "NFA_crop_account_value_per_ha_endpoint_ratio": (end_crop / b["area_crop_ha"]) / (start_crop / a["area_crop_ha"]),
        "noncrop_account_sum_start_gha_including_built_up": start_non_crop,
        "noncrop_account_sum_end_gha_including_built_up": end_non_crop,
        "noncrop_account_sum_endpoint_ratio": end_non_crop / start_non_crop,
        "noncrop_contribution_pp_of_B1961": 100 * (end_non_crop - start_non_crop) / a["B_gha"],
        "interpretation": "NFA account decomposition; per-ha quotient is not independent physical crop yield; non-cropland includes built-up land and is not ecological capital",
    }


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"No rows to write: {path}")
    fieldnames = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def make_fig2(rows: list[dict], metrics: dict) -> None:
    yrs = np.array([r["year"] for r in rows], dtype=float)
    B = np.array([r["B_gha"] for r in rows], dtype=float)
    Eprod = np.array([r["E_production_gha"] for r in rows], dtype=float)
    P = np.array([r["population_persons"] for r in rows], dtype=float)
    rprod = Eprod / B
    bpc = B / P
    eprod_pc = Eprod / P
    ix = lambda x: x / x[0]

    plt.rcParams.update({"font.size": 8, "axes.titlesize": 9, "axes.labelsize": 8})
    fig, ax = plt.subplots(1, 3, figsize=(11.7, 3.55), constrained_layout=True)
    ax[0].plot(yrs, ix(B), color="#2b6f8a", lw=2, label="Aggregate biocapacity $B$")
    ax[0].plot(yrs, ix(bpc), color="#d18f29", lw=2, label="Biocapacity per person $B/P$")
    ax[0].plot(yrs, ix(eprod_pc), color="#8b4a75", lw=2, label="Production footprint per person $E_{\\mathrm{prod}}/P$")
    ax[0].set_title("(a) Aggregate and per-capita indices")
    ax[0].set_ylabel("Index (1961 = 1)")
    ax[0].set_xlabel("Year")
    ax[0].grid(alpha=.25)
    ax[0].legend(frameon=False, fontsize=6.6, loc="best")

    ax[1].plot(yrs, rprod, color="#39764b", lw=2)
    ax[1].axhline(1, color="#555555", ls="--", lw=1)
    crossing = metrics["first_year_E_production_over_B_gt_1"]
    rc = rprod[np.where(yrs == crossing)[0][0]]
    ax[1].scatter([crossing], [rc], color="#bf4939", s=24, zorder=5)
    ax[1].annotate(f"first $E_{{\\mathrm{{prod}}}}/B>1$: {crossing}", (crossing, rc), xytext=(8, 9), textcoords="offset points", fontsize=6.8)
    ax[1].set_title("(b) Production footprint / biocapacity")
    ax[1].set_ylabel("$E_{\\mathrm{prod}}/B$ (dimensionless)")
    ax[1].set_xlabel("Year")
    ax[1].grid(alpha=.25)

    dln_p = metrics["dln_population"]
    dln_bpc = metrics["dln_B_percap"]
    dln_b = metrics["dln_B"]
    x = np.arange(3)
    ax[2].bar(x[0], dln_p, color="#2b6f8a", width=.64)
    ax[2].bar(x[1], dln_bpc, bottom=dln_p, color="#d18f29", width=.64)
    ax[2].bar(x[2], dln_b, color="#424b54", width=.64)
    ax[2].plot([x[0] + .32, x[1] + .32], [dln_p, dln_p], color="#666666", lw=.8, ls=":")
    ax[2].axhline(0, color="#555555", lw=.8)
    ax[2].set_xticks(x, [r"$\Delta\ln P$", r"$\Delta\ln(B/P)$", r"$\Delta\ln B$"])
    ax[2].set_title(f"(c) Exact log decomposition, 1961–{metrics['end_year']}")
    ax[2].set_ylabel("Log change (dimensionless)")
    ax[2].set_ylim(min(dln_bpc - .15, -0.9), dln_p + .2)
    ax[2].text(x[0], dln_p + .025, f"+{dln_p:.3f}", ha="center", va="bottom", fontsize=7)
    ax[2].text(x[1], dln_b + .025, f"{dln_bpc:.3f}", ha="center", va="bottom", fontsize=7)
    ax[2].text(x[2], dln_b + .025, f"+{dln_b:.3f}", ha="center", va="bottom", fontsize=7)
    ax[2].grid(axis="y", alpha=.25)
    fig.savefig(FIG2, dpi=300, bbox_inches="tight")
    plt.close(fig)


def make_fig3(components: list[dict], end_year: int) -> None:
    labels = [r["component"] for r in components]
    display_labels = ["Built-up land*" if label == "Built-up land" else label for label in labels]
    pp = np.array([r["contribution_pp_of_B1961"] for r in components])
    ratios = np.array([r["endpoint_ratio"] for r in components])
    colors = ["#2b6f8a" if value >= 0 else "#b54f3e" for value in pp]
    plt.rcParams.update({"font.size": 8, "axes.titlesize": 9, "axes.labelsize": 8})
    fig, ax = plt.subplots(1, 2, figsize=(10.9, 3.8))
    y = np.arange(len(labels))
    ax[0].barh(y, pp, color=colors, height=.62)
    ax[0].axvline(0, color="#555555", lw=.8)
    ax[0].set_yticks(y, display_labels)
    ax[0].invert_yaxis()
    ax[0].set_xlabel(f"Contribution to total biocapacity change, percentage points\n(1961–{end_year}; denominator $B_{{1961}}$)")
    ax[0].set_title("(a) Component account contributions")
    ax[0].grid(axis="x", alpha=.25)
    for yi, value in enumerate(pp):
        if value < -2:
            ax[0].text(value / 2, yi, f"{value:+.2f}", va="center", ha="center", fontsize=7, color="white", fontweight="bold")
        else:
            ax[0].text(value + (0.25 if value >= 0 else -0.25), yi, f"{value:+.2f}", va="center", ha="left" if value >= 0 else "right", fontsize=7)

    ax[1].barh(y, ratios, color="#638d78", height=.62)
    ax[1].axvline(1, color="#555555", ls="--", lw=1)
    ax[1].set_yticks(y, display_labels)
    ax[1].invert_yaxis()
    ax[1].set_xscale("log")
    ax[1].set_xlim(.65, 12.3)
    ax[1].set_xlabel(f"Endpoint NFA account ratio $B_i({end_year})/B_i(1961)$ (log scale)")
    ax[1].set_title("(b) Component account ratios")
    ax[1].grid(axis="x", alpha=.25, which="both")
    for yi, value in enumerate(ratios):
        ax[1].text(value * (1.04 if value >= 1 else .96), yi, f"{value:.2f}×", va="center", ha="left" if value >= 1 else "right", fontsize=7)

    fig.subplots_adjust(left=.205, right=.985, bottom=.23, top=.91, wspace=.24)
    fig.text(
        .205, .025,
        "* Built-up land is an NFA accounting class, not ecological capital. All bars are account values, not physical-stock measurements.",
        ha="left", va="bottom", fontsize=7, color="#333333",
    )
    fig.savefig(FIG3, dpi=300, bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    source = load_world()
    rows = [enrich(row) for row in source]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(DATA_DIR / "nfa_2026_world_1961_2023.csv", rows)

    endpoint_rows = [endpoint_metrics(rows, year) for year in (2022, 2023)]
    write_csv(DATA_DIR / "nfa_2026_endpoint_metrics.csv", endpoint_rows)
    component_rows = []
    for year in (2022, 2023):
        component_rows.extend(component_metrics(rows, year))
    write_csv(DATA_DIR / "nfa_2026_component_metrics.csv", component_rows)
    aggregate_component_rows = [aggregate_component_metrics(rows, year) for year in (2022, 2023)]
    write_csv(DATA_DIR / "nfa_2026_crop_and_noncrop_account_metrics.csv", aggregate_component_rows)

    # Validate the exact population identity and account-component sum (integer-gha rounding).
    for metrics in endpoint_rows:
        if abs(metrics["dln_identity_residual"]) > 1e-12:
            raise AssertionError(f"Population identity failed in {metrics['period']}: {metrics['dln_identity_residual']}")
        if abs(metrics["component_sum_residual_start_gha"]) > 2 or abs(metrics["component_sum_residual_end_gha"]) > 2:
            raise AssertionError(f"NFA components differ from the total by more than 2 gha in {metrics['period']}")
    for components in (component_rows[:len(COMPONENTS)], component_rows[len(COMPONENTS):]):
        if abs(sum(row["contribution_pp_of_B1961"] for row in components) - next(m["B_change_pct"] for m in endpoint_rows if m["period"] == components[0]["period"])) > 1e-6:
            raise AssertionError(f"Component contributions do not sum to total change in {components[0]['period']}")

    primary = next(metrics for metrics in endpoint_rows if metrics["end_year"] == 2023)
    primary_components = [row for row in component_rows if row["period"] == "1961-2023"]
    make_fig2(rows, primary)
    make_fig3(primary_components, 2023)

    summary = {
        "source_workbook": "data/NEFBA_2026.xlsx; official NFA Edition 2026, results and metadata v1.0",
        "source_sheet": "world_data",
        "source_measure": "efp_total_gha = global Ecological Footprint of Production; biocap_total_gha = biocapacity",
        "primary_period": "1961-2023 (complete published world observations; estimated 2024-2025 rows excluded)",
        "same_edition_sensitivity": "1961-2022, using the same workbook and population series",
        "population_source": POPULATION_SOURCE,
        "per_capita_method": "Recomputed from total account gha divided by pop_1000_persons * 1000; rounded workbook per-capita columns are not used",
        "component_status": "NFA biocapacity account values; not physical-stock measurements. Built-up land is not ecological capital.",
        "endpoint_metrics": endpoint_rows,
        "component_metrics": component_rows,
        "crop_and_noncrop_account_metrics": aggregate_component_rows,
        "figures": [str(FIG2.relative_to(ROOT)), str(FIG3.relative_to(ROOT))],
    }
    (DATA_DIR / "nfa_2026_summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")

    print("NFA 2026 endpoint audit (production footprint explicitly labeled)")
    for metrics in endpoint_rows:
        print(
            f"{metrics['period']}: B={metrics['B_change_pct']:.3f}% (Δln={metrics['dln_B']:.6f}); "
            f"B/P={metrics['B_percap_change_pct']:.3f}% (Δln={metrics['dln_B_percap']:.6f}); "
            f"P Δln={metrics['dln_population']:.6f}; "
            f"E_prod/P={metrics['E_production_percap_change_pct']:.3f}%; "
            f"first E_prod/B>1={metrics['first_year_E_production_over_B_gt_1']}"
        )
    print("Wrote:", DATA_DIR / "nfa_2026_world_1961_2023.csv")
    print("Wrote:", DATA_DIR / "nfa_2026_endpoint_metrics.csv")
    print("Wrote:", DATA_DIR / "nfa_2026_component_metrics.csv")
    print("Wrote:", DATA_DIR / "nfa_2026_crop_and_noncrop_account_metrics.csv")
    print("Wrote:", FIG2)
    print("Wrote:", FIG3)


if __name__ == "__main__":
    main()
