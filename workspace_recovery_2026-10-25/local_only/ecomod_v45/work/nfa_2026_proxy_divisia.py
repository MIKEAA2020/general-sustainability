#!/usr/bin/env python3
"""Recompute the exploratory crop-proxy Divisia sensitivity for ECOMOD v45.

This is NOT a physical-capital estimate. It creates one calibrated account
proxy X = b_hat_f * A_f and an accounting residual C = B_NFA - X. The independent
input series are FAOSTAT-derived cropland area and cereal yield distributed by
Our World in Data; the proxy's base-year level is anchored to the 1961 NFA
cropland-biocapacity account. The mapping between cereal yield, FAOSTAT
cropland area, NFA global hectares, and the model's fast book is an explicit
proxy assumption, not a measured identity. C is not identified as physical
capital or as the model quantity Upsilon(q) A_c.

Inputs (packaged):
  data/NEFBA_2026.xlsx                         NFA 2026 official workbook
  data/empirical/cropland_area_owid.csv        FAOSTAT cropland area via OWID
  data/empirical/cereal_yield_owid.csv         FAOSTAT cereal yield via OWID

Outputs:
  data/nfa_2026_proxy_input_series.csv
  data/nfa_2026_proxy_divisia.csv              measured 1961 NFA crop-share anchor
  data/nfa_2026_proxy_sensitivity.csv          alternative phi_0 calibrations

Run from any directory: python work/nfa_2026_proxy_divisia.py
Requires openpyxl only.
"""
from __future__ import annotations

import csv
import math
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NFA_FILE = ROOT / "data" / "NEFBA_2026.xlsx"
AREA_FILE = ROOT / "data" / "empirical" / "cropland_area_owid.csv"
YIELD_FILE = ROOT / "data" / "empirical" / "cereal_yield_owid.csv"
OUT_SERIES = ROOT / "data" / "nfa_2026_proxy_input_series.csv"
OUT_PRIMARY = ROOT / "data" / "nfa_2026_proxy_divisia.csv"
OUT_SENS = ROOT / "data" / "nfa_2026_proxy_sensitivity.csv"
END_YEARS = (2022, 2023)
PHI0_SENSITIVITY = (0.10, 0.12864722023747283, 0.15, 0.19, 0.25, 0.30, 0.32, 0.34)


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"No rows to write: {path}")
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def load_nfa_world() -> dict[int, dict]:
    if not NFA_FILE.is_file():
        raise FileNotFoundError(NFA_FILE)
    workbook = load_workbook(NFA_FILE, read_only=True, data_only=True)
    sheet = workbook["world_data"]
    headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    out = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if row.get("country_name") == "World" and row.get("year") is not None and 1961 <= int(row["year"]) <= 2023:
            out[int(row["year"])] = {
                "quality_flag": row["flag"],
                "B_gha": float(row["biocap_total_gha"]),
                "B_crop_NFA_gha": float(row["biocap_crop_gha"]),
                "population_1000_persons": float(row["pop_1000_persons"]),
            }
    if not set(range(1961, 2024)).issubset(out):
        raise ValueError("NFA 2026 world_data must contain 1961--2023")
    if any(out[year]["quality_flag"] == "E" for year in range(1961, 2024)):
        raise ValueError("Estimated row found inside the complete 1961--2023 NFA window")
    return out


def load_owid_world(path: Path, value_column: str, value_name: str) -> dict[int, float]:
    if not path.is_file():
        raise FileNotFoundError(path)
    values: dict[int, float] = {}
    with path.open(newline="", encoding="utf-8-sig") as file:
        for row in csv.DictReader(file):
            if row.get("Code") != "OWID_WRL":
                continue
            year = int(row["Year"])
            value = row.get(value_column)
            if value not in (None, ""):
                values[year] = float(value)
    needed = set(range(1961, 2024))
    missing = sorted(needed - set(values))
    if missing:
        raise ValueError(f"{value_name}: missing World observations for {missing}")
    return values


def logmean(a: float, b: float) -> float:
    if a <= 0 or b <= 0:
        raise ValueError("Logarithmic mean requires positive values")
    if math.isclose(a, b, rel_tol=1e-13, abs_tol=0):
        return (a + b) / 2
    return (b - a) / math.log(b / a)


def decompose(nfa: dict[int, dict], area: dict[int, float], cereal_yield: dict[int, float], end_year: int, phi0: float) -> dict:
    start_year = 1961
    n0, nt = nfa[start_year], nfa[end_year]
    a0, at = area[start_year], area[end_year]
    y0, yt = cereal_yield[start_year], cereal_yield[end_year]
    b0, bt = n0["B_gha"], nt["B_gha"]

    # X0 = phi0 * B0; area and physical cereal-yield index provide independent
    # growth factors; b_hat_f is calibrated so X0 is in gha.
    x0 = phi0 * b0
    dln_area = math.log(at / a0)
    dln_yield = math.log(yt / y0)
    x_ratio = (at / a0) * (yt / y0)
    xt = x0 * x_ratio
    c0 = b0 - x0
    ct = bt - xt
    dln_b = math.log(bt / b0)
    row = {
        "period": f"{start_year}-{end_year}",
        "start_year": start_year,
        "end_year": end_year,
        "phi0_base_share": phi0,
        "phi0_is_2026_NFA_cropland_share": abs(phi0 - n0["B_crop_NFA_gha"] / b0) < 1e-12,
        "B_start_gha": b0,
        "B_end_gha": bt,
        "NFA_cropland_account_start_gha": n0["B_crop_NFA_gha"],
        "NFA_cropland_account_end_gha": nt["B_crop_NFA_gha"],
        "NFA_cropland_share_1961": n0["B_crop_NFA_gha"] / b0,
        "OWID_FAOSTAT_cropland_area_start_ha": a0,
        "OWID_FAOSTAT_cropland_area_end_ha": at,
        "OWID_FAOSTAT_cropland_area_ratio": at / a0,
        "OWID_FAOSTAT_cropland_area_dln": dln_area,
        "OWID_FAOSTAT_cereal_yield_start_t_ha": y0,
        "OWID_FAOSTAT_cereal_yield_end_t_ha": yt,
        "OWID_FAOSTAT_cereal_yield_ratio": yt / y0,
        "OWID_FAOSTAT_cereal_yield_dln": dln_yield,
        "proxy_fast_account_X_start_gha": x0,
        "proxy_fast_account_X_end_gha": xt,
        "proxy_fast_account_X_endpoint_ratio": x_ratio,
        "accounting_residual_C_start_gha": c0,
        "accounting_residual_C_end_gha": ct,
        "residual_C_positive_both_endpoints": c0 > 0 and ct > 0,
        "dln_B": dln_b,
        "proxy_fast_share_start": x0 / b0,
        "proxy_fast_share_end": xt / bt,
        "interpretation": "Illustrative account proxy only; C is not measured physical capital and is not identified as Upsilon(q)A_c",
    }
    if c0 > 0 and ct > 0:
        wx = logmean(x0, xt) / logmean(b0, bt)
        wc = logmean(c0, ct) / logmean(b0, bt)
        dln_c = math.log(ct / c0)
        yield_contribution = wx * dln_yield
        area_contribution = wx * dln_area
        residual_contribution = wc * dln_c
        row.update({
            "log_decomposition_valid": True,
            "dln_C": dln_c,
            "w_X": wx,
            "w_C": wc,
            "yield_contribution_to_dln_B": yield_contribution,
            "area_contribution_to_dln_B": area_contribution,
            "residual_contribution_to_dln_B": residual_contribution,
            "sum_of_contributions": yield_contribution + area_contribution + residual_contribution,
            "identity_residual": dln_b - yield_contribution - area_contribution - residual_contribution,
            "level_delta_X_gha": xt - x0,
            "level_delta_C_gha": ct - c0,
        })
    else:
        delta_x = xt - x0
        delta_c = (bt - b0) - delta_x
        row.update({
            "log_decomposition_valid": False,
            "dln_C": "",
            "w_X": "",
            "w_C": "",
            "yield_contribution_to_dln_B": "",
            "area_contribution_to_dln_B": "",
            "residual_contribution_to_dln_B": "",
            "sum_of_contributions": "",
            "identity_residual": "",
            "level_delta_X_gha": delta_x,
            "level_delta_C_gha": delta_c,
        })
    return row


def main() -> None:
    nfa = load_nfa_world()
    area = load_owid_world(AREA_FILE, "Cropland - Area (hectares)", "cropland area")
    cereal_yield = load_owid_world(YIELD_FILE, "Cereals - Yield (tonnes per hectare)", "cereal yield")

    series_rows = []
    for year in range(1961, 2024):
        series_rows.append({
            "year": year,
            "NFA_2026_quality_flag": nfa[year]["quality_flag"],
            "NFA_2026_B_gha": nfa[year]["B_gha"],
            "NFA_2026_cropland_biocapacity_gha": nfa[year]["B_crop_NFA_gha"],
            "NFA_2026_cropland_share_of_B": nfa[year]["B_crop_NFA_gha"] / nfa[year]["B_gha"],
            "OWID_FAOSTAT_cropland_area_ha": area[year],
            "OWID_FAOSTAT_cereal_yield_t_ha": cereal_yield[year],
        })
    write_csv(OUT_SERIES, series_rows)

    phi0_measured = nfa[1961]["B_crop_NFA_gha"] / nfa[1961]["B_gha"]
    primary_rows = [decompose(nfa, area, cereal_yield, year, phi0_measured) for year in END_YEARS]
    for row in primary_rows:
        if row["log_decomposition_valid"] and abs(row["identity_residual"]) > 1e-12:
            raise AssertionError(f"Proxy Divisia identity failed in {row['period']}: {row['identity_residual']}")
    write_csv(OUT_PRIMARY, primary_rows)

    sens_rows = []
    for end_year in END_YEARS:
        for phi0 in PHI0_SENSITIVITY:
            sens_rows.append(decompose(nfa, area, cereal_yield, end_year, phi0))
    for row in sens_rows:
        if row["log_decomposition_valid"] and abs(row["identity_residual"]) > 1e-12:
            raise AssertionError(f"Sensitivity identity failed in {row['period']}, phi0={row['phi0_base_share']}")
    write_csv(OUT_SENS, sens_rows)

    print("NFA 2026 + OWID/FAOSTAT proxy audit (proxy residual is not physical capital)")
    for row in primary_rows:
        if row["log_decomposition_valid"]:
            print(
                f"{row['period']} (phi0={row['phi0_base_share']:.6f}): "
                f"dlnB={row['dln_B']:.6f}; dln(area)={row['OWID_FAOSTAT_cropland_area_dln']:.6f}; "
                f"dln(cereal yield)={row['OWID_FAOSTAT_cereal_yield_dln']:.6f}; "
                f"Divisia yield={row['yield_contribution_to_dln_B']:.6f}, "
                f"area={row['area_contribution_to_dln_B']:.6f}, "
                f"residual={row['residual_contribution_to_dln_B']:.6f}; "
                f"C_end={row['accounting_residual_C_end_gha']:.3e} gha"
            )
        else:
            print(f"{row['period']}: log decomposition invalid; use level decomposition")
    print("Wrote:", OUT_SERIES)
    print("Wrote:", OUT_PRIMARY)
    print("Wrote:", OUT_SENS)


if __name__ == "__main__":
    main()
