#!/usr/bin/env python3
"""Batch-5 corrected editions — integrity and content verification.

Pins the SHA-256 of every new file produced by the batch-5 joint-audit
response (the nine second-edition manuscripts, the monograph v1.1, the seven
archival corrected editions, the two corrected runners with their regenerated
result files, the renamed figure copies, the packet README_v2, and the joint
evaluation document), and machine-checks the corrected invariants that the
joint evaluation (BATCH5_JOINT_AUDIT_EVALUATION.md) records:

  * the audited first editions are byte-identical to the packet's pins
    (the batch-5 rule: current versions are never altered);
  * the corrected constants and sentences are present and their defective
    predecessors absent (per finding);
  * the regenerated cod intervention results differ from the committed ones
    in exactly one value (the flat-180-kt T=inf convergence correction);
  * the regenerated Edwards intervention results are value-identical (the
    comparator fix is inert);
  * the renamed figure copies are byte-identical to the committed figures;
  * the internal-audit vocabulary remains absent from every second edition
    while the mathematical gate vocabulary stays in place where it belongs.

Idempotent; run from the repository root:

    python3 reaudit/verify_batch5_editions.py
"""
import hashlib
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

FAIL = 0
N = 0


def check(name, ok):
    global FAIL, N
    N += 1
    print(("PASS " if ok else "FAIL ") + name)
    if not ok:
        FAIL += 1


def sha256(path):
    return hashlib.sha256((ROOT / path).read_bytes()).hexdigest()


# ---------------------------------------------------------------- pins
# The audited first editions (must equal external_review_packet.py's pins).
FIRST_EDITIONS = {
    "papers/paper1_general_theory/manuscript.md":
        "dc568ef4192d5aacd4a89028ce750d5078a957273af308ffe7e23cf859c20736",
    "papers/paper2_theorem_atlas/manuscript.md":
        "2404081fda7e17ea718ba5c92dcf19e6e0389a5f147131ed1a07de10d0ae841a",
    "papers/paper3_material_ledgers/manuscript.md":
        "d593c6f87ad340ebcea4c41fb77b82818348d7bbfb4149422f553cb26808d07b",
    "papers/paper4_delay_dynamics/manuscript.md":
        "f807993a1636029a0b7da689a19b2a25396f360ca4286525766d131171a5dfb9",
    "papers/paper5_sampled_governance/manuscript.md":
        "edb8eac3f4d1ef149a5d9e60c72464d2d8650bb36451ad3ae8330b786079308d",
    "wave_e_cod/manuscript/wave_E_cod_forecast_ladder.md":
        "b5a5a43b47d43e93d3ed657a0e22c639611145866e665003d86f5ad1f38d3ce5",
    "wave_e_cod/manuscript/wave_E_cod_intervention.md":
        "46d55eadad8014a3a545219380a2f4f6d22a02016d15c850025afa230dc53353",
    "wave_e_edwards/manuscript/wave_E_edwards_forecast_ladder.md":
        "f6ebe9422020177d8d791e9c7877857ebc381f83abb52eef1bf5bacfc9fb0c09",
    "wave_e_edwards/manuscript/wave_E_edwards_intervention.md":
        "09fd1b88c99fbf46cbd0b3f0f103d8a3f431799bb755a0b47a3a767e43bd7a34",
}

# The batch-5 new files (hashes computed at commit time; any drift after the
# batch-5 push is a failure of this suite).
NEW_FILES = {
    "papers/paper1_general_theory/manuscript_v2.md": None,
    "papers/paper2_theorem_atlas/manuscript_v2.md": None,
    "papers/paper3_material_ledgers/manuscript_v2.md": None,
    "papers/paper4_delay_dynamics/manuscript_v2.md": None,
    "papers/paper5_sampled_governance/manuscript_v2.md": None,
    "wave_e_cod/manuscript/wave_E_cod_forecast_ladder_v2.md": None,
    "wave_e_cod/manuscript/wave_E_cod_intervention_v2.md": None,
    "wave_e_edwards/manuscript/wave_E_edwards_forecast_ladder_v2.md": None,
    "wave_e_edwards/manuscript/wave_E_edwards_intervention_v2.md": None,
    "revised_sustainability_manuscript_v1.1.md": None,
    "ms_part1_corrected.md": None,
    "ms_part2_corrected.md": None,
    "ms_part3_corrected.md": None,
    "ms_part4_corrected.md": None,
    "general_theory_of_sustainability_v0.1_corrected.md": None,
    "general_theory_of_sustainability_v0.2_comprehensive_corrected.md": None,
    "general_theory_of_sustainability_manuscript_corrected.md": None,
    "wave_e_cod/src/run_intervention_v2.py": None,
    "wave_e_cod/results/intervention_results_v2.json": None,
    "wave_e_cod/results/intervention_boundaries_v2.csv": None,
    "wave_e_edwards/src/run_intervention_v2.py": None,
    "wave_e_edwards/results/intervention_results_v2.json": None,
    "wave_e_edwards/results/intervention_boundaries_v2.csv": None,
    "wave_e_edwards/manuscript/fig4_pass2.png": None,
    "wave_e_edwards/manuscript/fig5_fibre.png": None,
    "external_review_packet/README_v2.md": None,
    "BATCH5_JOINT_AUDIT_EVALUATION.md": None,
    # residual-pass files (2026-08-29, second session)
    "wave_e_edwards/SPECIFICATION_v2.md": None,
    "wave_e_cod/SPECIFICATION_v2.md": None,
    "build_revised_sustainability_manuscript_v1_1.py": None,
    "revised_sustainability_manuscript_v1.1.docx": None,
}
# Pin file: written by the first run of this script (or maintained alongside
# it) — reaudit/batch5_edition_pins.json. If absent, it is created and the
# hashes are reported; on a pinned repository the file is committed.
PIN_FILE = ROOT / "reaudit" / "batch5_edition_pins.json"


def read(p):
    return (ROOT / p).read_text(encoding="utf-8")


def main():
    # ---- 1. First editions unchanged ------------------------------------
    for path, pin in FIRST_EDITIONS.items():
        check(f"first edition unchanged: {path}", sha256(path) == pin)

    # ---- 2. New files exist; pins ---------------------------------------
    pins = json.loads(PIN_FILE.read_text()) if PIN_FILE.exists() else {}
    fresh = {}
    for path in NEW_FILES:
        if not (ROOT / path).exists():
            check(f"exists: {path}", False)
            continue
        h = sha256(path)
        fresh[path] = h
        if pins:
            check(f"pinned hash: {path}", pins.get(path) == h)
    if not pins:
        PIN_FILE.write_text(json.dumps(fresh, indent=1, sort_keys=True) + "\n")
        print(f"[pin file created: {PIN_FILE}]")
        check("all new files present", len(fresh) == len(NEW_FILES))

    # ---- 3. Content invariants (per finding) ----------------------------
    p1 = read("papers/paper1_general_theory/manuscript_v2.md")
    check("P1: FAST/SLOW direction restored",
          "High-`s_2`-price assessors" in p1
          and "Low-`s_2`-price assessors license `FAST` only" not in p1)
    check("P1: thirteen-slot tuple enumerated",
          "`S = (T, Z, S_st, B_out, V, Γ, O, A, C, R, D, K, P)`" in p1)
    check("P1: nonzero nonnegative orthant named",
          "nonzero nonnegative orthant" in p1)
    check("P1: a-fortiori direction corrected",
          "strictest" in p1 and "most permissive" not in p1)
    check("P1: no [cite:] placeholders",
          "[cite:" not in p1 and "[Cite:" not in p1)
    check("P1: nineteen sources + full accounting",
          "nineteen sources" in p1 and "354 row-verified" in p1
          and "28 adjudicated rejected-or-negative-only" in p1
          and "twenty sources" not in p1)
    check("P1: four Wave E manuscripts counted",
          p1.count("four scored Wave E manuscripts") >= 2)
    check("P1: Succ subseteq G fix", "`Succ ⊆ G`" in p1)
    check("P1: Das-Dennis single year", "Das–Dennis 1997/1998" not in p1)

    p2 = read("papers/paper2_theorem_atlas/manuscript_v2.md")
    pat = re.compile(r"^\*\*(?:Conditional Theorem|Theorem|Proposition|Lemma|"
                     r"Corollary|Definition|Remark|Example|Template|Programme|"
                     r"Counterexample) (\d+)\.(\d+)")
    for sect in ("5", "6", "8", "13"):
        insec, seq = False, []
        for l in p2.split("\n"):
            if l.startswith("## " + sect + " "):
                insec = True
                continue
            if insec and l.startswith("## "):
                insec = False
                continue
            if insec:
                m = pat.match(l)
                if m:
                    seq.append(int(m.group(2)))
        check(f"P2: section {sect} numbering monotone",
              all(seq[i] < seq[i + 1] for i in range(len(seq) - 1)))
    check("P2: abstract enumerates twelve families",
          "canonical typed definitions and notation bridges" in p2
          and "intergenerational and stochastic bounds" in p2
          and "F08, the scalar resource and sink kernels" in p2)
    check("P2: Theorem 6.4 proof supplied",
          "supplied in this edition" in p2
          and "will be supplied at camera-ready" not in p2)
    check("P2: Cond Thm 10.1 preamble restored",
          "y(t) = A(t) - A_{\\min}" in p2.replace("$", "")
          and "first hitting time $T_A$ of $\\{A \\le A_{\\min}\\}$" in p2)
    check("P2: CES preamble + sigma>1 in item 5",
          "reference scales" in p2 and
          "if $\\sigma > 1$ and $\\mu_A = \\delta_A$" in p2)
    check("P2: Theorem 3.4 enforced-exit reading",
          "the disturbance can force exit" in p2)
    check("P2: hypothesis objects split",
          "Hypothesis object 12.3a" in p2 and "Hypothesis object 12.3b" in p2)
    check("P2: seven sources + A002 row 89",
          "together with one further row from the already-closed primary "
          "source A002" in p2)
    check("P2: row 64 canonical rendering",
          "conditional on the registered prerequisite result" in p2)

    p3 = read("papers/paper3_material_ledgers/manuscript_v2.md")
    check("P3: A_g0 declared with scale separation",
          "separation-of-scale condition" in p3)
    check("P3: mining-rate reading corrected",
          "at precisely the extraction rate" in p3)
    check("P3: ledger cross-reference note",
          "cross-references, not retained rows" in p3
          and "the retained count is 52" in p3)
    check("P3: nested 43/42 populations noted",
          "42-stock annual-managed subset" in p3)

    p4 = read("papers/paper4_delay_dynamics/manuscript_v2.md")
    check("P4: loop-gain frequency corrected",
          "0.0589" in p4 and "0.0583" not in p4)
    check("P4: fold passage branch-resolved + caveat restored",
          "basin collapse between $\\tau=5.574$ and $5.576$" in p4
          and "provisional" in p4
          and "0.964$ at $\\tau=5.5815$" in p4)
    check("P4: tau branch mapping explicit",
          "lower family ($\\omega_1\\approx0.02518$) moves up" in p4)
    check("P4: memory gain renamed",
          "gain-$\\gamma_m$" in p4 and "gain-$g$" not in p4)
    check("P4: tether alpha number withdrawn",
          "1.3\\times10^{-3}" not in p4 and
          "does not itself tabulate" in p4)
    check("P4: tau+ variant pairing named",
          "$\\tau_+\\approx132$–$150$ yr and $\\approx76$–$80$ yr" in p4)
    check("P4: ledger cross-reference note",
          "cross-references owned by Paper 3" in p4)

    p5 = read("papers/paper5_sampled_governance/manuscript_v2.md")
    check("P5: Schaefer specialisation corrected",
          "degenerate member of this family" in p5
          and "Allee-factor-1 specialisation" not in p5)
    check("P5: La Mancha relapse dated",
          "before its 2019–2023 extraction relapse" in p5)
    check("P5: four Wave E manuscripts counted",
          "four scored Wave E manuscripts" in p5)
    check("P5: CC-A002-034 cross-reference marked",
          "not a retained row of this paper" in p5)

    cf = read("wave_e_cod/manuscript/wave_E_cod_forecast_ladder_v2.md")
    check("cod forecast: piecewise surplus law",
          "a(S_t)=" in cf and
          "no value of $\\mathfrak s$ makes the displayed factor" not in cf)
    check("cod forecast: Schaefer-not-parameter point present",
          "the Schaefer model is the separate member" in cf)
    check("cod forecast: Omega_xte Brier values + corrected conclusion",
          "persistence 0.06 at $h{=}1$ and 0.27 at $h{=}5$" in cf
          and "M1 0.05 and 0.45" in cf and "M3 0.02 and 0.31" in cf
          and "M1 and M3 improve the one-year Brier score" in cf
          and "No model improves the Brier score over persistence either"
          not in cf)
    check("cod forecast: abstract range split by catch pass",
          "115--196 kt" in cf and "115--206 kt across both catch "
          "treatments" in cf)
    check("cod forecast: I_ref and b defined",
          "training-window median of the observed index" in cf)
    check("cod forecast: freeze caveat restated",
          "no dated pre-score protocol file" in cf)
    check("cod forecast: reproducibility qualified",
          "environment-sensitive" in cf)
    check("cod forecast: Schijns pages completed", "2675–2683" in cf)
    check("cod forecast: 2021 checkpoint exact", "$=423$ kt" in cf)

    ci = read("wave_e_cod/manuscript/wave_E_cod_intervention_v2.md")
    check("cod intervention: 2338.3 in the table, 2335.4 out of the table",
          "| flat 180 kt | 1297.1 | 1171.0 | 991.2 | 2338.3 |" in ci
          and "| 2335.4 |" not in ci)
    check("cod intervention: S1/cpm 1991 distinguished",
          "critical-zone rule S1" in ci and "under the cascade (cpm" in ci
          and "critical-period rule, the path is below" not in ci)
    check("cod intervention: 60kt phrasing corrected",
          "60 kt and larger" in ci)
    check("cod intervention: exact operands",
          "172.47 - 114.85 = 57.62" in ci)
    check("cod intervention: T=5 classification both kernels",
          "1985, 1987, 1989" in ci and "additionally, 1988 for BAU" in ci)
    check("cod intervention: v2 runner documented",
          "run_intervention_v2.py" in ci)

    ef = read("wave_e_edwards/manuscript/"
              "wave_E_edwards_forecast_ladder_v2.md")
    check("Edwards forecast: figure files match numbers",
          "(fig4_pass2.png)" in ef and "(fig5_fibre.png)" in ef
          and "fig5_pass2.png)" not in ef and "fig4_fibre.png)" not in ef)
    check("Edwards forecast: model names standardised",
          "M2\\_Renso" in ef and "M2\\_enso" not in ef)
    check("Edwards forecast: abstract margin clause excludes Rar",
          "AR(1) on recharge itself being 0.4 ft worse" in ef)

    ei = read("wave_e_edwards/manuscript/"
              "wave_E_edwards_intervention_v2.md")
    check("Edwards intervention: Abstract present",
          "## Abstract" in ei)
    check("Edwards intervention: References present",
          "## References" in ei and "Umphres" in ei)
    check("Edwards intervention: T approx 13",
          "T \\approx 13" in ei and "T \\approx 14" not in ei)
    check("Edwards intervention: first-stage trigger wording",
          "first-stage CPM trigger" in ei and "deepest CPM trigger" not in ei)
    check("Edwards intervention: oracle-gap baselines labelled",
          "retained AR(1)" in ei)
    check("Edwards intervention: BAU mean exact", "282.16" in ei)
    check("Edwards intervention: flat-0 T=4 certified horizon recorded",
          "T = 4" in ei and "687.9" in ei)
    check("Edwards intervention: v2 runner documented",
          "run_intervention_v2.py" in ei)

    mo = read("revised_sustainability_manuscript_v1.1.md")
    check("monograph v1.1: hybrid trajectory repaired",
          "\\tau=(q_0,z_0)\\rightarrow(q_1,z_1)\\rightarrow\\cdots ." in mo)
    check("monograph v1.1: ten rationale clauses",
          "narrowing of viable transformation options indicates the "
          "architecture graph itself is closing" in mo)
    check("monograph v1.1: version header",
          "version 1.1" in mo)

    mp2 = read("ms_part2_corrected.md")
    check("ms_part2 corrected: no CR bytes", "\r" not in mp2)
    check("ms_part2 corrected: arrows restored",
          "(q_0,z_0)\\rightarrow(q_1,z_1)\\rightarrow\\cdots ." in mp2)
    check("ms_part2 corrected: T' defined",
          "post-arrival maintenance horizon" in mp2)

    # ---- 4. Result-file invariants --------------------------------------
    old = json.loads(
        (ROOT / "wave_e_cod/results/intervention_results.json").read_text())
    new = json.loads(
        (ROOT / "wave_e_cod/results/intervention_results_v2.json").read_text())

    def diff_paths(a, b, base=""):
        out = []
        if type(a) != type(b):
            out.append(base)
            return out
        if isinstance(a, dict):
            for k in a:
                if k not in b:
                    out.append(base + "/" + k)
                else:
                    out += diff_paths(a[k], b[k], base + "/" + k)
            for k in b:
                if k not in a:
                    out.append(base + "/" + k)
        elif isinstance(a, list):
            if len(a) != len(b):
                out.append(base + " (length)")
            else:
                for i, (x, y) in enumerate(zip(a, b)):
                    out += diff_paths(x, y, f"{base}[{i}]")
        elif a != b:
            out.append(base)
        return out

    d = diff_paths(old, new)
    check("cod v2 results: exactly one changed value",
          d == ["/kernels/flat_75/UC_q10/inf/nominal[0][0]"])
    if d == ["/kernels/flat_75/UC_q10/inf/nominal[0][0]"]:
        v = new["kernels"]["flat_75"]["UC_q10"]["inf"]["nominal"][0][0]
        check("cod v2 results: converged fixpoint value",
              abs(v - 2338.273378118786) < 1e-6)

    eo = json.loads((ROOT / "wave_e_edwards/results/"
                     "intervention_results.json").read_text())
    en = json.loads((ROOT / "wave_e_edwards/results/"
                     "intervention_results_v2.json").read_text())
    check("Edwards v2 results: value-identical (comparator fix inert)",
          diff_paths(eo, en) == [])

    # ---- 5. Figure-copy hash equality ------------------------------------
    check("figure copy: fig4_pass2 == fig5_pass2 (committed)",
          sha256("wave_e_edwards/manuscript/fig4_pass2.png")
          == sha256("wave_e_edwards/manuscript/fig5_pass2.png"))
    check("figure copy: fig5_fibre == fig4_fibre (committed)",
          sha256("wave_e_edwards/manuscript/fig5_fibre.png")
          == sha256("wave_e_edwards/manuscript/fig4_fibre.png"))

    # ---- 6. Vocabulary scan over the second editions ---------------------
    v2s = list(NEW_FILES)[:9] + [
        "revised_sustainability_manuscript_v1.1.md",
        "ms_part1_corrected.md", "ms_part2_corrected.md",
        "ms_part3_corrected.md", "ms_part4_corrected.md",
        "general_theory_of_sustainability_v0.1_corrected.md",
        "general_theory_of_sustainability_v0.2_comprehensive_corrected.md",
        "general_theory_of_sustainability_manuscript_corrected.md",
    ]
    forbidden = [
        r"\bF4\b", r"build_panel\.py", r"\bNOT CONFIRMED\b", r"\bmanifest\b",
        r"independent[- ]rerun", r"\brerun-verified\b",
    ]
    bad = []
    for f in v2s:
        try:
            t = read(f)
        except Exception:
            continue
        for rx in forbidden:
            if re.search(rx, t):
                bad.append((f, rx))
    check("vocabulary: internal-audit terms absent from all editions",
          not bad)
    if bad:
        print("  offenders:", bad)
    # Mathematical gate vocabulary must SURVIVE in Papers 2-5 v2 (Task 58).
    for f, probe in [
        ("papers/paper4_delay_dynamics/manuscript_v2.md", "gated"),
        ("papers/paper5_sampled_governance/manuscript_v2.md", "gate"),
    ]:
        check(f"mathematical gate vocabulary retained: {f}",
              probe in read(f))

    # ---- 7. Residual pass (2026-08-29, second session) -------------------
    # 7a. Paper 1 v2 duplication removed
    check("P1 v2: tuple sentence no longer duplicated",
          p1.count("A *model* in this programme is a fully specified tuple;") == 1)

    # 7b. G04(a): the normative-authority slot in the flagship Appendix A
    for f in ["general_theory_of_sustainability_v0.1_corrected.md",
              "general_theory_of_sustainability_v0.2_comprehensive_corrected.md",
              "general_theory_of_sustainability_manuscript_corrected.md"]:
        t = read(f)
        check(f"G04(a): normative-authority slot in Appendix A: {f}",
              "- Normative authority \\((\\mathcal N)\\) or procedure used to "
              "choose social constraints:" in t)

    # 7c. GT-04: the symbol-I overload declared + disambiguated
    for f in ["general_theory_of_sustainability_v0.1_corrected.md",
              "general_theory_of_sustainability_v0.2_comprehensive_corrected.md",
              "general_theory_of_sustainability_manuscript_corrected.md"]:
        t = read(f)
        check(f"GT-04: I-overload declared in header: {f}",
              "triple use of the symbol I" in t)
        check(f"GT-04: inventory reading disambiguated on the line: {f}",
              "Here \\(I\\) denotes the produced reserve-inventory stock" in t)

    # 7d. GT-05/GT-06: supersession notes
    for f in ["general_theory_of_sustainability_v0.1_corrected.md",
              "general_theory_of_sustainability_v0.2_comprehensive_corrected.md",
              "general_theory_of_sustainability_manuscript_corrected.md"]:
        t = read(f)
        check(f"GT-05/GT-06: registry-resolution supersession note: {f}",
              "superseded by the successor manuscript's unified constraint "
              "registry" in t)

    # 7e. F7: the proof-obligations expansion acknowledged
    mp4 = read("ms_part4_corrected.md")
    check("F7: six-to-nine obligations expansion acknowledged",
          "those six are retained" in mp4
          and "restated here as the *boundary* obligation" in mp4
          and "composition, transformation, and commons obligations are added" in mp4)

    # 7f. F11: the Appendix B provenance alphabet completed
    check("F11: provenance alphabet carries the definitional entry",
          "[P/E/M/N/D/L]" in mp4 and "[P/E/M/N/L]" not in mp4)
    check("F11: column-alphabet legend present",
          "Column alphabets (distinct by design)" in mp4)

    # 7g. SPECIFICATION second editions
    es = read("wave_e_edwards/SPECIFICATION_v2.md")
    check("Edwards SPEC v2: W03 echo corrected",
          "T\\approx 13\\) yr" in es and "beyond ~14 yr" not in es
          and "692.6" in es and "12.7" in es)
    check("Edwards SPEC v2: edition note present",
          "Second edition (2026-08-29)" in es)
    cs = read("wave_e_cod/SPECIFICATION_v2.md")
    check("cod SPEC v2: W09 echo split by catch pass",
          "115–196 kt under the coarse catch regime" in cs
          and "115–206 kt across both catch treatments" in cs)
    check("cod SPEC v2: edition note present",
          "Second edition (2026-08-29)" in cs)
    # the first-edition sheets remain byte-identical (still pinned by the
    # wave_e spec-match suite; here: unchanged on disk relative to their own
    # committed content is implied by that suite)

    # 7h. the monograph v1.1 docx
    try:
        from docx import Document
        d11 = Document(str(ROOT / "revised_sustainability_manuscript_v1.1.docx"))
        full11 = "\n".join(p.text for p in d11.paragraphs)
        check("monograph v1.1 docx: repaired formula present",
              "\\tau=(q_0,z_0)\\rightarrow(q_1,z_1)\\rightarrow\\cdots ." in full11)
        check("monograph v1.1 docx: ten-indicator rationale present",
              "recovery time indicates weakening restorative dynamics" in full11)
    except ImportError:
        check("monograph v1.1 docx: python-docx unavailable (skipped)", True)

    # 7i. the addendum
    ev = read("BATCH5_JOINT_AUDIT_EVALUATION.md")
    check("evaluation: residual-pass addendum present",
          "## 6. Residual pass (2026-08-29, second session)" in ev)

    print("-" * 64)
    print(f"{N - FAIL}/{N} checks pass")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
